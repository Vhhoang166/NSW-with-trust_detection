from __future__ import annotations

import csv
import html
import math
import os
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw, ImageFont


os.environ.setdefault("XDG_CACHE_HOME", "/tmp")
os.environ.setdefault("MPLCONFIGDIR", "/tmp/matplotlib")

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SUMMARY_CSV = ROOT / "output_400_rounds" / "all_approaches_400rounds_summary.csv"
COMPARISON_CSV = ROOT / "output_400_rounds" / "all_approaches_200_vs_400_ratio_comparison.csv"
RUN_NOTES = ROOT / "output_400_rounds" / "NSW_ALL_APPROACHES_400_RUN_NOTES.txt"
LEGACY_NOTES = ROOT / "outputs" / "Output_Moradi" / "NSW_RUN_SUMMARY.txt"
FULL_4500_DIR = DOCS / "combined_full_4500_summary"
FULL_4500_CELL_CSV = FULL_4500_DIR / "full_4500_cell_summary.csv"
FULL_4500_TRIAL_CSV = FULL_4500_DIR / "full_4500_trial_summary.csv"
FULL_4500_README = FULL_4500_DIR / "README.txt"
FULL_4500_FIGURE_DIR = FULL_4500_DIR / "figures"
HTML_OUT = DOCS / "trust_aware_online_nsw_manuscript.html"
DOCX_OUT = DOCS / "trust_aware_online_nsw_report.docx"


APPROACH_LABELS = {
    "baseline_fixed_alpha": "Baseline fixed-alpha set-aside",
    "adaptive_alpha": "Adaptive trust-aware alpha",
    "pace_trusted": "PACE-inspired trusted allocator",
    "generalized_mean_greedy": "Generalized-mean greedy",
    "sample_resolving": "Sample-resolving approximation",
    "expert_advice": "Expert-advice controller",
    "robust_aggregation": "Robust aggregation",
    "pure_greedy_alpha0": "Pure greedy alpha=0 reference",
    "pure_equal_alpha1": "Pure equal alpha=1 reference",
}

ATTACK_LABELS = {
    "byzantine": "Byzantine",
    "burst": "Burst",
    "reputation_poison": "Reputation poison",
    "trust_mimicry": "Trust mimicry",
    "compound": "Compound",
}

TRUST_MODE_LABELS = {
    "algorithm1": "Algorithm 1 hard gate",
    "no_trust": "No trust",
    "oracle": "Oracle trust",
    "sparse_online": "Sparse-online gate",
}

APPROACH_DESCRIPTIONS = {
    "baseline_fixed_alpha": (
        "Original trust-filtered set-aside plus greedy water-filling with a fixed "
        "allocation_alpha. This is the control condition."
    ),
    "adaptive_alpha": (
        "Keeps the original allocator but replaces fixed alpha with alpha_t, "
        "raising the defensive set-aside when trust or report instability rises."
    ),
    "pace_trusted": (
        "Uses a PACE-inspired allocation rule on the trusted set. It tests whether "
        "pacing-style state improves online welfare under the same trust filter."
    ),
    "generalized_mean_greedy": (
        "Replaces the Nash-specific greedy marginal rule with a generalized-mean "
        "marginal rule, linking the simulator to generalized-welfare online allocation."
    ),
    "sample_resolving": (
        "Stabilizes current reports with a rolling historical median and then uses "
        "a generalized-mean allocation step. It is an intentionally lightweight "
        "approximation of full sample-based re-solving."
    ),
    "expert_advice": (
        "Runs defensive set-aside, aggressive set-aside, and generalized-mean "
        "candidates, then updates exponential weights from realized performance."
    ),
    "robust_aggregation": (
        "Clips reports using median absolute deviation before passing them into "
        "the baseline allocation rule."
    ),
    "pure_greedy_alpha0": (
        "A reference extreme with no set-aside protection. It is not the proposed "
        "method; it shows the high-efficiency edge of the alpha tradeoff."
    ),
    "pure_equal_alpha1": (
        "A reference extreme with all budget split equally across the trusted set. "
        "It is included to anchor the defensive end of the alpha tradeoff."
    ),
}

APPROACH_GUIDE = [
    {
        "key": "baseline_fixed_alpha",
        "role": "Clean baseline allocator",
        "intuition": (
            "This is the reference rule that the rest of the report modifies. After the "
            "trust gate decides which agents are currently trusted, the allocator splits "
            "the unit resource into two conceptual parts: a protective equal-share part "
            "and a greedy welfare-improving part. The fixed alpha parameter controls how "
            "much weight is put on the protective side."
        ),
        "implementation": (
            "In the simulator, alpha is fixed for the entire run. The method therefore "
            "cannot react differently to a quiet period, an obvious attack, or a stealth "
            "attack that passes the detector. Its purpose is to provide a stable, directly "
            "interpretable baseline for the allocation-side comparisons."
        ),
        "interpretation": (
            "If a new variant only slightly improves over this baseline, the improvement "
            "should be interpreted carefully. The fixed-alpha rule is simple, but it is "
            "also intentionally conservative and easy to audit."
        ),
    },
    {
        "key": "adaptive_alpha",
        "role": "Proposed direct extension",
        "intuition": (
            "Adaptive alpha keeps the same set-aside structure but makes the defense "
            "level time-dependent. The core idea is that the allocator should not use the "
            "same risk posture when reports look stable as when trust evidence or reports "
            "become unstable."
        ),
        "implementation": (
            "The simulator computes an alpha_t signal from trust/report instability. When "
            "instability rises, alpha_t moves toward a more defensive set-aside behavior; "
            "when the environment appears calmer, it can allow more welfare-seeking "
            "greedy allocation. This makes the method a controller plus allocator rather "
            "than a completely new allocation theory."
        ),
        "interpretation": (
            "This is the main direct project contribution because it minimally changes "
            "the original baseline. It is not claimed to win every setting; its value is that "
            "it connects the trust layer to the allocation risk level in a transparent way. "
            "Unlike expert advice, adaptive alpha does not combine multiple allocators; "
            "it modifies only the alpha choice inside the original set-aside structure."
        ),
    },
    {
        "key": "pace_trusted",
        "role": "Paper-inspired allocator approximation",
        "intuition": (
            "PACE-style allocation can be understood as trying to pace resources according "
            "to agents' accumulated utility state. Instead of only asking who has the "
            "largest immediate report, the method uses a state variable that reflects how "
            "allocation has accumulated over time."
        ),
        "implementation": (
            "This report uses a trusted-set approximation of that idea: the rule is applied "
            "after the trust filter and is simplified to fit the same simulator interface "
            "as the other variants. It should not be read as a full reproduction of the "
            "source paper's model or guarantees."
        ),
        "interpretation": (
            "The method is useful as a benchmark because it asks whether pacing-style "
            "state helps under strategic reports. A weak result here is evidence about "
            "this simplified implementation, not a rejection of the original PACE theory."
        ),
    },
    {
        "key": "generalized_mean_greedy",
        "role": "Alternative welfare-shape allocator",
        "intuition": (
            "Nash social welfare is one member of a broader family of welfare objectives. "
            "A generalized-mean rule changes the marginal priority assigned to agents, "
            "which can shift the allocation toward different fairness-efficiency tradeoffs."
        ),
        "implementation": (
            "The simulator implements a greedy marginal rule inspired by generalized-mean "
            "welfare and applies it to the trusted set. The evaluation metric, however, "
            "remains Nash social welfare ratio. This means the method is being tested as "
            "an allocation variant inside the shared benchmark, not under every assumption "
            "of the source paper."
        ),
        "interpretation": (
            "This variant is included to see whether changing the welfare shape helps when "
            "reports are attacked. If it does not beat fixed-alpha set-aside on NSW ratio, "
            "that may simply mean the altered objective is not aligned with this metric "
            "and trust environment."
        ),
    },
    {
        "key": "sample_resolving",
        "role": "Historical-sample approximation",
        "intuition": (
            "Sample-based re-solving methods use information from previous observations "
            "to make better online allocation decisions. The hope is that historical "
            "structure can stabilize decisions that would otherwise overreact to one "
            "round of noisy or adversarial reports."
        ),
        "implementation": (
            "The implemented version is deliberately lightweight: it uses rolling report "
            "medians as a proxy for a richer sample-based re-solving algorithm, then "
            "feeds the stabilized signal into an allocation step. This keeps the runtime "
            "manageable and compatible with the existing simulator."
        ),
        "interpretation": (
            "The main risk is that history itself can be poisoned. If stealth attackers "
            "survive the trust gate, their reports can enter the rolling history, so "
            "stabilization may preserve contamination rather than remove it."
        ),
    },
    {
        "key": "expert_advice",
        "role": "Meta-controller over allocator candidates",
        "intuition": (
            "Expert advice is not a standalone allocator. It is a portfolio method: run "
            "several candidate allocation behaviors, observe which ones perform better, "
            "and gradually put more weight on the stronger candidates."
        ),
        "implementation": (
            "In this simulator, the experts are allocation candidates such as defensive "
            "set-aside, aggressive set-aside, and generalized-mean behavior. The controller "
            "updates exponential weights from realized performance and then combines the "
            "candidate decisions."
        ),
        "interpretation": (
            "This explains why expert advice can rank highly: it is allowed to adapt among "
            "several behaviors while fixed-alpha set-aside commits to one behavior. The "
            "result should be reported as an empirical controller benchmark, not as a "
            "claim that the 1997 expert-advice theorem directly solves this allocation problem."
        ),
    },
    {
        "key": "robust_aggregation",
        "role": "Preprocessing heuristic",
        "intuition": (
            "Robust aggregation tries to clean the input before allocation. Rather than "
            "changing the allocator itself, it clips suspiciously extreme reports so that "
            "one very large or very small value has less influence."
        ),
        "implementation": (
            "The implementation uses a median/MAD-style clipping rule and then passes the "
            "adjusted reports into the baseline allocation rule. It is therefore a "
            "preprocessing layer, not a new welfare optimizer."
        ),
        "interpretation": (
            "This is naturally suited to obvious outliers such as Byzantine manipulation. "
            "It is less likely to solve stealth attacks where malicious reports are designed "
            "to look statistically ordinary."
        ),
    },
    {
        "key": "pure_greedy_alpha0",
        "role": "Efficiency reference endpoint",
        "intuition": (
            "Pure greedy removes the protective set-aside and gives the allocation rule "
            "maximum freedom to chase immediate welfare gains among the trusted agents."
        ),
        "implementation": (
            "It corresponds to alpha=0 in the set-aside family. Under a good trust gate, "
            "this can produce a high NSW ratio because the allocator does not spend budget "
            "on equal-share protection."
        ),
        "interpretation": (
            "It is not the proposed robust method. Its high ratio is informative, but it "
            "also leaks the most resource to malicious agents under Algorithm 1, showing "
            "that it relies heavily on trust quality."
        ),
    },
    {
        "key": "pure_equal_alpha1",
        "role": "Defensive reference endpoint",
        "intuition": (
            "Pure equal allocation is the opposite endpoint. It ignores reported value "
            "differences within the trusted set and divides the resource equally."
        ),
        "implementation": (
            "It corresponds to alpha=1 in the set-aside family. This gives the largest "
            "protective floor but removes most of the welfare-seeking behavior."
        ),
        "interpretation": (
            "The method is included as an anchor for the alpha tradeoff. It is useful for "
            "understanding the defensive boundary, but it is too inefficient to be the "
            "main recommendation."
        ),
    },
]

CITATIONS = [
    {
        "id": "R1",
        "label": "Banerjee et al. 2022",
        "text": (
            "Siddhartha Banerjee, Vasilis Gkatzelis, Artur Gorokh, and Billy Jin. "
            "Online Nash Social Welfare Maximization with Predictions. SODA 2022, "
            "pages 1-19. DOI: 10.1137/1.9781611977073.1; arXiv:2008.03564."
        ),
        "url": "https://doi.org/10.1137/1.9781611977073.1",
    },
    {
        "id": "R2",
        "label": "Akgun et al. 2025",
        "text": (
            "Orhan Eren Akgun, Sarper Aydin, Stephanie Gil, and Angelia Nedic. "
            "Multi-Agent Trustworthy Consensus under Random Dynamic Attacks. "
            "arXiv:2504.07189, 2025."
        ),
        "url": "https://arxiv.org/abs/2504.07189",
    },
    {
        "id": "R3",
        "label": "Moradi et al. 2015",
        "text": (
            "Parham Moradi, Sajad Ahmadian, and Fardin Akhlaghian. An effective "
            "trust-based recommendation method using a novel graph clustering "
            "algorithm. Physica A 436:462-481, 2015. DOI: 10.1016/j.physa.2015.05.008."
        ),
        "url": "https://doi.org/10.1016/j.physa.2015.05.008",
    },
    {
        "id": "R4",
        "label": "Yang, Liao, and Kroer 2024",
        "text": (
            "Zongjun Yang, Luofeng Liao, and Christian Kroer. Greedy-Based Online "
            "Fair Allocation with Adversarial Input: Enabling Best-of-Many-Worlds "
            "Guarantees. AAAI 2024; arXiv:2308.09277."
        ),
        "url": "https://arxiv.org/abs/2308.09277",
    },
    {
        "id": "R5",
        "label": "Yang, Kumar, and Kroer 2026",
        "text": (
            "Zongjun Yang, Rachitesh Kumar, and Christian Kroer. Online "
            "Generalized-mean Welfare Maximization: Achieving Near-Optimal Regret "
            "from Samples. arXiv:2602.10469, 2026."
        ),
        "url": "https://arxiv.org/abs/2602.10469",
    },
    {
        "id": "R6",
        "label": "Freund and Schapire 1997",
        "text": (
            "Yoav Freund and Robert E. Schapire. A Decision-Theoretic Generalization "
            "of On-Line Learning and an Application to Boosting. Journal of Computer "
            "and System Sciences 55(1):119-139, 1997. DOI: 10.1006/jcss.1997.1504."
        ),
        "url": "https://doi.org/10.1006/jcss.1997.1504",
    },
    {
        "id": "R7",
        "label": "Frank and Wolfe 1956",
        "text": (
            "Marguerite Frank and Philip Wolfe. An algorithm for quadratic "
            "programming. Naval Research Logistics Quarterly 3(1-2):95-110, 1956. "
            "DOI: 10.1002/nav.3800030109."
        ),
        "url": "https://doi.org/10.1002/nav.3800030109",
    },
    {
        "id": "R8",
        "label": "von Luxburg 2007",
        "text": (
            "Ulrike von Luxburg. A tutorial on spectral clustering. Statistics and "
            "Computing 17:395-416, 2007. DOI: 10.1007/s11222-007-9033-z."
        ),
        "url": "https://doi.org/10.1007/s11222-007-9033-z",
    },
    {
        "id": "R9",
        "label": "Huber and Ronchetti 2009",
        "text": (
            "Peter J. Huber and Elvezio M. Ronchetti. Robust Statistics, 2nd edition. "
            "Wiley, 2009. Used as background for median/MAD-style robust clipping."
        ),
        "url": "https://doi.org/10.1002/9780470434697",
    },
]

SOURCE_MAP = [
    [
        "Baseline fixed-alpha set-aside allocator",
        "Banerjee-style online NSW with predictions: set-aside plus greedy/water-filling over predicted monopolist utilities.",
        "R1",
        "Direct baseline adaptation. The project restricts the allocator to the current trusted set; that trust gate is our simulator adaptation.",
    ],
    [
        "Prediction stress tests and V_tilde idea",
        "Prediction-based online NSW framework where predictions are monopolist-utility estimates.",
        "R1",
        "The named prediction scenarios are simulator stress tests, not separate source-paper methods.",
    ],
    [
        "Cumulative beta trust detector and xi_t threshold",
        "Trust-observation detector for trustworthy consensus under random dynamic attacks.",
        "R2",
        "The simulator uses the detector as a pre-allocation trust filter rather than as a consensus-only mechanism.",
    ],
    [
        "Moradi blended trust graph and sparsest-subgraph diagnostic",
        "Trust-aware recommender graph clustering with sparsest-subgraph initialization.",
        "R3",
        "The code adapts the trust graph idea for malicious-agent diagnostics in this allocation simulator.",
    ],
    [
        "Spectral detector comparison",
        "Standard spectral clustering baseline.",
        "R8",
        "Used only as a comparison detector, not as the main trust rule.",
    ],
    [
        "PACE-inspired trusted allocator",
        "PACE, Pacing According to Current Estimated utility, and its relation to greedy allocation under adversarial input.",
        "R4",
        "The implementation is a trusted-set approximation, not a full reproduction of all assumptions and guarantees.",
    ],
    [
        "Generalized-mean greedy",
        "Online generalized-mean welfare maximization with greedy rules.",
        "R5",
        "The project implements a lightweight greedy variant inside the same trust/attack simulator.",
    ],
    [
        "Sample-resolving approximation",
        "Re-solving paradigm using historical samples for generalized-mean welfare.",
        "R5",
        "The code uses rolling report medians as a cheap proxy, so this is only an approximation of the paper's re-solving methods.",
    ],
    [
        "Expert-advice controller",
        "Exponential weights / multiplicative-weights online learning.",
        "R6",
        "The experts are allocator candidates designed for this project: defensive set-aside, aggressive set-aside, and generalized-mean greedy.",
    ],
    [
        "Robust aggregation",
        "Robust-statistics idea behind median/MAD clipping.",
        "R9",
        "This is a project heuristic for clipping reports before the baseline allocator.",
    ],
    [
        "Frank-Wolfe offline benchmark",
        "Conditional-gradient / Frank-Wolfe optimization.",
        "R7",
        "Used to approximate the offline NSW denominator faster and more reliably than an SLSQP solve with equal-split fallback.",
    ],
    [
        "Byzantine, Burst, Reputation Poison, Trust Mimicry, Compound attacks",
        "Simulation attack scenarios.",
        "Project-defined",
        "These are not direct methods from the cited allocation papers; they are adversarial scenarios used to stress the trust-aware allocator.",
    ],
    [
        "Adaptive trust-aware alpha",
        "Project contribution inspired by the robustness-versus-efficiency role of alpha in the set-aside allocator.",
        "R1 plus project-defined controller",
        "Proposed controller. The dynamic alpha_t controller is ours; it is not claimed as a theorem from Banerjee et al.",
    ],
]

APPROACH_TAXONOMY = [
    [
        "Fixed-alpha set-aside",
        "Allocator",
        "Mostly direct adaptation",
        "Clean baseline",
    ],
    [
        "Adaptive alpha",
        "Controller plus allocator",
        "Project-defined",
        "Proposed method",
    ],
    [
        "PACE-inspired",
        "Allocator approximation",
        "Paper-inspired approximation",
        "Baseline variant",
    ],
    [
        "Generalized-mean greedy",
        "Allocator variant",
        "Paper-inspired approximation",
        "Baseline variant",
    ],
    [
        "Sample-resolving",
        "Approximation",
        "Paper-inspired approximation",
        "Stress-test baseline",
    ],
    [
        "Expert advice",
        "Meta-controller",
        "Meta-controller",
        "Empirical upper-style controller",
    ],
    [
        "Robust aggregation",
        "Preprocessing heuristic",
        "Heuristic",
        "Robustness baseline",
    ],
    [
        "Pure greedy alpha=0",
        "Reference extreme",
        "Reference endpoint",
        "Efficiency-side alpha anchor",
    ],
    [
        "Pure equal alpha=1",
        "Reference extreme",
        "Reference endpoint",
        "Defense-side alpha anchor",
    ],
]

CONTRIBUTION_TABLE = [
    [
        "Allocation rule",
        "Fixed-alpha set-aside greedy.",
        "Adaptive trust-aware alpha that changes the set-aside level over time.",
    ],
    [
        "Use of trust",
        "Trust is mainly a hard gate: agents are included or excluded.",
        "Trust instability also controls allocation risk through alpha_t.",
    ],
    [
        "Robustness level",
        "One fixed defensive level for every attack regime.",
        "Dynamic defense level that can become more conservative when reports or trust become unstable.",
    ],
    [
        "Evaluation focus",
        "Standard attack comparisons around the fixed baseline.",
        "Shared benchmark with stealth attacks and multiple implementation-level variants.",
    ],
]

FINAL_RECOMMENDATION_TABLE = [
    ["Use as baseline", "Fixed-alpha set-aside", "Cleanest directly sourced allocation baseline."],
    ["Use as proposed method", "Adaptive alpha", "Cleanest allocation-side extension; its advantage depends on trust quality."],
    ["Use as empirical controller", "Expert advice", "Strongest practical controller among non-extreme variants."],
    ["Use as efficiency reference", "Pure greedy alpha=0", "Highest Algorithm 1 ratio, but high malicious resource share makes it a reference endpoint."],
    ["Use as defensive reference", "Pure equal alpha=1", "Defense-side alpha endpoint; useful anchor but low NSW efficiency."],
    ["Use as trust upper bound", "Oracle trust", "Shows allocation performance when malicious identities are known."],
    ["Use as promising trust extension", "Sparse-online gate", "Strong trust-gate ablation result, not the main allocation contribution."],
]

PRIMARY_4500_FIGURES = [
    (
        FULL_4500_FIGURE_DIR / "fig1_trust_mode_ratio.png",
        "Figure 1. Mean NSW ratio by trust mode in the 4,500-trial benchmark.",
    ),
    (
        FULL_4500_FIGURE_DIR / "fig2_malicious_share_by_trust_mode.png",
        "Figure 2. Mean malicious resource share by trust mode.",
    ),
    (
        FULL_4500_FIGURE_DIR / "fig3_algorithm1_approach_ranking.png",
        "Figure 3. Allocation approach ranking under the Algorithm 1 hard trust gate.",
    ),
    (
        FULL_4500_FIGURE_DIR / "fig4_attack_trust_heatmap.png",
        "Figure 4. Attack-by-trust-mode NSW ratio heatmap.",
    ),
    (
        FULL_4500_FIGURE_DIR / "fig5_trust_gate_comparison_by_attack.png",
        "Figure 5. Algorithm 1, sparse-online, and oracle trust comparison by attack type.",
    ),
]

MISSING_EXPERIMENTS = [
    [
        "Critical",
        "Trust ablation: no trust, Algorithm 1 hard gate, sparse hard gate, soft trust, oracle trust",
        "Separate detector effects from allocator effects.",
        "NSW ratio, minimum utility, malicious resource share, FPR/FNR.",
    ],
    [
        "Critical",
        "Sparse detector used online",
        "Test whether high sparse F1 improves welfare when it actually controls allocation.",
        "Algorithm 1 gate vs sparse gate vs blended/soft gate.",
    ],
    [
        "Critical",
        "Stealth-attack allocation/resource leakage",
        "Show how much resource reaches malicious agents when Reputation Poison or Trust Mimicry evades detection.",
        "Allocation heatmaps and malicious resource share for stealth attacks.",
    ],
    [
        "Critical",
        "Adaptive-alpha trajectory",
        "Show that alpha_t rises when trust/report instability rises.",
        "alpha_t, trust instability, NSW ratio, and minimum utility over time.",
    ],
    [
        "Critical",
        "Expert-advice weight trajectory",
        "Verify that the empirical win comes from switching among experts rather than a static mixture.",
        "Expert weights over time by attack type.",
    ],
    [
        "Important",
        "Compound attack across all seven approaches",
        "Evaluate the coordinated attack type already present in the Moradi final code.",
        "Ratio ranking and detector behavior under Compound.",
    ],
    [
        "Important",
        "Paired significance or bootstrap confidence intervals",
        "Make adaptive-alpha and expert-advice lifts more statistically credible.",
        "Paired seed lift with 95% confidence intervals.",
    ],
]

FIGURE_STRATEGY = [
    [
        "All-approaches ratio heatmap",
        "Main story",
        "Overall empirical ranking under the shared Algorithm 1 benchmark.",
    ],
    [
        "Detector behavior table / misclassification plots",
        "Main story",
        "Shows why trust alone is insufficient: Algorithm 1 catches obvious attacks but fails on stealth attacks.",
    ],
    [
        "Adaptive-alpha approach-specific signals",
        "Main story",
        "Mechanism evidence showing how the proposed method changes its defensive level over time.",
    ],
    [
        "Trust-mode ablation and malicious resource share",
        "Main story",
        "Separates allocator behavior from detector behavior and quantifies resource leakage.",
    ],
    [
        "Legacy Moradi focused-run plots",
        "Appendix / supporting context",
        "Useful for background, but not part of the primary 4,500-trial benchmark.",
    ],
    [
        "Full output graph set",
        "Grouped appendix",
        "Useful for auditability, but should not replace the main narrative.",
    ],
]

ORIGINAL_PROJECT_DIRECTIONS = [
    [
        "Alpha selection",
        "Where should alpha come from, and how should the robustness-efficiency tradeoff be chosen?",
        "Addressed through fixed-alpha baseline and the adaptive-alpha proposed method.",
    ],
    [
        "Trust mechanism",
        "Can malicious agents be detected or scored more reliably?",
        "Used mostly as a fixed experimental environment in this report; detector diagnostics explain input contamination.",
    ],
    [
        "Allocation alternatives",
        "Under the same trust-aware setting, which allocation rules reduce welfare loss when attackers pass the filter?",
        "Main focus of this report.",
    ],
]

TRUST_MODE_OVERVIEW = [
    [
        "No trust",
        "Everyone remains eligible for allocation.",
        "Lower-bound / no-defense condition: shows what happens when allocation receives contaminated input directly.",
    ],
    [
        "Algorithm 1 hard gate",
        "The beta-gap trust detector decides which agents are excluded before allocation.",
        "Main setting: evaluates allocation variants under the current trust mechanism.",
    ],
    [
        "Oracle trust",
        "The simulator removes exactly the malicious agents before allocation.",
        "Upper-bound trust-quality condition: separates allocation weakness from detector weakness.",
    ],
    [
        "Sparse-online gate",
        "A graph-based Moradi-style sparse detector is used online as the trust gate.",
        "Diagnostic trust extension: tests how allocation changes when trust quality improves, but it uses a strong malicious-count prior.",
    ],
]

SOURCE_TO_SIMULATOR_MAPPING = [
    [
        "Fixed-alpha set-aside",
        "Reserve an equal set-aside share for fairness, then use greedy/water-filling behavior for efficiency.",
        "Online divisible goods with prediction information, especially monopolist-utility style predictions.",
        "Run the set-aside allocator only on the trusted set returned by the trust gate.",
        "Tests whether the original robustness-efficiency structure remains stable when malicious agents may be admitted or excluded by trust.",
    ],
    [
        "Adaptive alpha",
        "Use the alpha tradeoff in set-aside allocation as a controllable robustness level.",
        "Not a source-paper theorem; it is a project-defined controller inspired by the set-aside baseline.",
        "Replace fixed alpha with alpha_t computed from trust/report instability.",
        "A direct allocation-side contribution: it modifies only the baseline alpha choice, not the whole allocator family.",
    ],
    [
        "PACE-inspired",
        "Pace allocation using current estimated utility state rather than only immediate reports.",
        "PACE-style guarantees depend on specific online fair-allocation assumptions.",
        "Use a simplified trusted-set pacing score inside the same trust-filtered simulator.",
        "Useful as a behavioral benchmark; weak performance is about this simplified implementation, not the original theorem.",
    ],
    [
        "Generalized-mean greedy",
        "Optimize a broader family of welfare objectives beyond Nash social welfare.",
        "The source setting studies generalized-mean welfare with sampling/regret structure.",
        "Implement a lightweight generalized-mean marginal rule after trust filtering.",
        "Tests whether broader welfare shaping helps when the evaluation metric is still NSW under contaminated reports.",
    ],
    [
        "Sample-resolving",
        "Use samples or history to re-solve online allocation decisions more reliably.",
        "Full sample-based methods assume richer sample/re-solving structure than this simulator uses.",
        "Use rolling report medians as a cheap historical proxy before allocation.",
        "Tests whether history stabilizes allocation; it can fail when stealth attackers poison the history.",
    ],
    [
        "Expert advice",
        "Use multiplicative weights to combine experts according to realized performance.",
        "Online learning over expert rewards/losses, not a direct NSW allocation algorithm by itself.",
        "Treat allocator candidates as experts and update weights over defensive, aggressive, and generalized behaviors.",
        "A meta-controller benchmark: strong empirical performance should not be read as an old theorem beating a newer allocator theorem.",
    ],
    [
        "Robust aggregation",
        "Clip or downweight outlying inputs using robust-statistics ideas.",
        "Robust statistics gives input-stabilization tools rather than a full online allocation rule.",
        "Clip reports with median/MAD-style logic before passing them to the baseline allocator.",
        "Useful for obvious outliers; less targeted to stealth attacks that look statistically normal.",
    ],
]

KEY_FINDINGS = [
    [
        "Allocation-side adaptation helps.",
        "Under Algorithm 1, expert advice and adaptive alpha improve over fixed-alpha set-aside among the main non-extreme variants.",
    ],
    [
        "Trust quality dominates under stealth attacks.",
        "Algorithm 1 fails on Reputation Poison and Trust Mimicry, so the allocator often receives contaminated trusted input.",
    ],
    [
        "Sparse-online is promising but not final.",
        "It nearly matches oracle trust in this simulator, but it uses a strong malicious-count prior and needs sensitivity testing.",
    ],
    [
        "Pure greedy is not the final answer.",
        "It achieves the highest NSW ratio under Algorithm 1, but it leaks the most resource to malicious agents and relies heavily on trust quality.",
    ],
    [
        "The next direction is allocation-trust co-design.",
        "The evidence points toward soft trust weighting, better allocation risk signals, and robustness tests for sparse-online trust.",
    ],
]

ALLOCATION_ALGORITHM_COMPARISON = [
    [
        "Fixed-alpha set-aside",
        "Yes, baseline allocator",
        "Hard gate: only trusted agents receive allocation.",
        "Keeps alpha fixed.",
        "Simple, interpretable, directly sourced baseline.",
        "Cannot adapt when trust becomes uncertain or stealth attackers pass the gate.",
    ],
    [
        "Adaptive alpha",
        "Controller plus allocator",
        "Trust/report instability controls alpha_t.",
        "Makes defensive set-aside dynamic.",
        "Cleanest project contribution; minimal extension of baseline.",
        "Mechanism depends on the calibration of the instability signal.",
    ],
    [
        "PACE-inspired",
        "Allocator approximation",
        "Runs on the trusted set.",
        "Tests pacing-style allocation under trust filtering.",
        "Useful comparator for online fair allocation behavior.",
        "Not a full reproduction of the PACE theorem setting.",
    ],
    [
        "Generalized-mean greedy",
        "Allocator variant",
        "Runs on the trusted set.",
        "Changes the welfare/marginal-priority shape.",
        "Can explore broader fairness objectives.",
        "May not improve NSW when NSW remains the evaluation metric.",
    ],
    [
        "Sample-resolving",
        "Historical-sample approximation",
        "Uses report history after trust filtering.",
        "Stabilizes current reports with historical samples.",
        "Useful negative-result stress baseline.",
        "History can be poisoned when stealth attackers survive trust filtering.",
    ],
    [
        "Expert advice",
        "Meta-controller",
        "Combines allocator candidates under the same trust environment.",
        "Shifts weights among defensive/aggressive/generalized behaviors.",
        "Strongest empirical controller benchmark.",
        "Mechanism attribution is less direct because it mixes candidate allocators.",
    ],
    [
        "Robust aggregation",
        "Preprocessing heuristic",
        "Clips reports before allocation.",
        "Reduces extreme reported values.",
        "Useful for extreme Byzantine-style manipulation.",
        "Less naturally targeted to stealth mimicry or reputation poisoning.",
    ],
    [
        "Pure greedy alpha=0",
        "Reference allocator extreme",
        "Runs on the selected trust set.",
        "Removes the set-aside floor entirely.",
        "Shows the high-efficiency edge of the alpha tradeoff.",
        "Sends more resource to malicious agents when trust fails; not a clean robustness contribution.",
    ],
    [
        "Pure equal alpha=1",
        "Reference allocator extreme",
        "Runs on the selected trust set.",
        "Uses only equal split within the trusted set.",
        "Shows the most defensive alpha endpoint.",
        "Sacrifices NSW efficiency and is not competitive as the main method.",
    ],
]

PROBLEM_FORMULATION = [
    [
        "Resource arrival",
        "At each round t = 1,...,T, one divisible resource with budget 1 arrives.",
    ],
    [
        "True value",
        "Each agent i has true value v_{i,t} for the round-t resource.",
    ],
    [
        "Report observed by algorithm",
        "The algorithm observes report r_{i,t}; for malicious agents this may differ from v_{i,t}.",
    ],
    [
        "Allocation decision",
        "Choose x_{i,t} >= 0 with sum_i x_{i,t} = 1, usually after applying the trust filter.",
    ],
    [
        "Cumulative utility",
        "U_{i,T} = sum_{t=1}^T v_{i,t} x_{i,t}. Evaluation uses true values, not reports.",
    ],
    [
        "Nash social welfare",
        "NSW_T = (prod_{i in legitimate} U_{i,T})^{1/n_legit}; equivalently exp((1/n_legit) sum_i log U_{i,T}).",
    ],
    [
        "NSW ratio",
        "Ratio = NSW_online / NSW_offline, where NSW_offline is the clairvoyant offline benchmark for the same valuation sequence.",
    ],
]


def read_csv(path: Path) -> list[dict]:
    with path.open(newline="") as f:
        rows = list(csv.DictReader(f))
    for row in rows:
        for key, value in list(row.items()):
            if key not in {"approach", "attack_type", "trust_mode"}:
                row[key] = float(value) if value not in {"", None} else float("nan")
    return rows


def normalize_full_4500_rows(rows: list[dict]) -> list[dict]:
    normalized = []
    for row in rows:
        r = dict(row)
        r["n_trials_completed"] = r.get("n_trials", 25.0)
        r["n_rounds"] = 400.0
        r["avg_ratio"] = r["final_nsw_ratio_mean"]
        r["std_ratio"] = r.get("final_nsw_ratio_std", float("nan"))
        r["avg_nsw"] = r["final_nsw_legit_mean"]
        r["avg_min_util"] = r["final_min_util_mean"]
        r["avg_fairness_gap"] = r["final_fairness_gap_mean"]
        r["avg_dr"] = r["final_detection_rate_mean"]
        r["avg_fpr"] = r["final_fp_rate_mean"]
        r["avg_malicious_resource_share"] = r["avg_malicious_resource_share_mean"]
        r["avg_offline_nsw"] = (
            r["final_nsw_legit_mean"] / r["final_nsw_ratio_mean"]
            if r["final_nsw_ratio_mean"] else float("nan")
        )
        normalized.append(r)
    return normalized


def group_by(rows: list[dict], key: str) -> dict[str, list[dict]]:
    grouped: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        grouped[row[key]].append(row)
    return dict(grouped)


def mean(values: list[float]) -> float:
    clean = [v for v in values if isinstance(v, (int, float)) and v == v]
    return sum(clean) / len(clean) if clean else float("nan")


def sample_sd(values: list[float]) -> float:
    clean = [v for v in values if isinstance(v, (int, float)) and v == v]
    if len(clean) < 2:
        return 0.0
    mu = mean(clean)
    return math.sqrt(sum((v - mu) ** 2 for v in clean) / (len(clean) - 1))


def ci95(values: list[float]) -> float:
    clean = [v for v in values if isinstance(v, (int, float)) and v == v]
    if len(clean) < 2:
        return 0.0
    return 1.96 * sample_sd(clean) / math.sqrt(len(clean))


def fmt_ratio(value: float) -> str:
    return f"{value:.3f}"


def fmt_pct(value: float) -> str:
    return f"{100.0 * value:.1f}%"


def fmt_float(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def rel(path: Path) -> str:
    return html.escape(os.path.relpath(path, DOCS).replace(os.sep, "/"))


def display_path(path: Path) -> str:
    resolved = path.resolve()
    docs_resolved = DOCS.resolve()
    try:
        if resolved.is_relative_to(docs_resolved):
            return resolved.relative_to(docs_resolved).as_posix()
    except AttributeError:
        try:
            return resolved.relative_to(docs_resolved).as_posix()
        except ValueError:
            pass
    try:
        return path.relative_to(ROOT).as_posix()
    except ValueError:
        return path.as_posix()


def paired_difference_rows(trial_rows: list[dict]) -> list[dict]:
    if not trial_rows:
        return []

    def key_for(row: dict, include_approach: bool = False) -> tuple:
        key = (row["attack_type"], int(row["trial"]))
        if include_approach:
            key = (row["approach"],) + key
        return key

    def paired_lift(
        label: str,
        left_filter,
        right_filter,
        include_approach: bool = False,
    ) -> dict:
        left = {key_for(r, include_approach): r for r in trial_rows if left_filter(r)}
        right = {key_for(r, include_approach): r for r in trial_rows if right_filter(r)}
        diffs = [
            left[k]["final_nsw_ratio"] - right[k]["final_nsw_ratio"]
            for k in sorted(set(left) & set(right))
        ]
        return {
            "comparison": label,
            "mean_diff": mean(diffs),
            "ci95": ci95(diffs),
            "n_pairs": len(diffs),
        }

    return [
        paired_lift(
            "Expert advice vs fixed-alpha under Algorithm 1",
            lambda r: r["trust_mode"] == "algorithm1" and r["approach"] == "expert_advice",
            lambda r: r["trust_mode"] == "algorithm1" and r["approach"] == "baseline_fixed_alpha",
        ),
        paired_lift(
            "Adaptive alpha vs fixed-alpha under Algorithm 1",
            lambda r: r["trust_mode"] == "algorithm1" and r["approach"] == "adaptive_alpha",
            lambda r: r["trust_mode"] == "algorithm1" and r["approach"] == "baseline_fixed_alpha",
        ),
        paired_lift(
            "Sparse-online trust vs Algorithm 1 hard gate",
            lambda r: r["trust_mode"] == "sparse_online",
            lambda r: r["trust_mode"] == "algorithm1",
            include_approach=True,
        ),
    ]


def compute_tables(
    summary_rows: list[dict],
    comparison_rows: list[dict],
    trial_rows: list[dict] | None = None,
) -> dict:
    algorithm1_rows = [r for r in summary_rows if r.get("trust_mode") == "algorithm1"]
    if not algorithm1_rows:
        algorithm1_rows = summary_rows

    by_app = group_by(algorithm1_rows, "approach")
    by_attack = group_by(algorithm1_rows, "attack_type")
    app_rankings = []
    for app, rows in by_app.items():
        app_rankings.append(
            {
                "approach": app,
                "mean_ratio": mean([r["avg_ratio"] for r in rows]),
                "mean_nsw": mean([r["avg_nsw"] for r in rows]),
                "mean_min_util": mean([r["avg_min_util"] for r in rows]),
                "mean_fairness_gap": mean([r["avg_fairness_gap"] for r in rows]),
                "mean_malicious_share": mean([r.get("avg_malicious_resource_share", float("nan")) for r in rows]),
                "mean_dr": mean([r["avg_dr"] for r in rows]),
            }
        )
    app_rankings.sort(key=lambda r: r["mean_ratio"], reverse=True)

    best_by_attack = {}
    for attack, rows in by_attack.items():
        best_by_attack[attack] = max(rows, key=lambda r: r["avg_ratio"])

    baseline = {
        r["attack_type"]: r
        for r in algorithm1_rows
        if r["approach"] == "baseline_fixed_alpha"
    }
    lifts = []
    for app, rows in by_app.items():
        if app == "baseline_fixed_alpha":
            continue
        abs_lifts = [
            r["avg_ratio"] - baseline[r["attack_type"]]["avg_ratio"]
            for r in rows
            if r["attack_type"] in baseline
        ]
        rel_lifts = [
            r["avg_ratio"] / baseline[r["attack_type"]]["avg_ratio"] - 1.0
            for r in rows
            if r["attack_type"] in baseline
        ]
        lifts.append(
            {
                "approach": app,
                "abs_lift": mean(abs_lifts),
                "rel_lift": mean(rel_lifts),
            }
        )
    lifts.sort(key=lambda r: r["abs_lift"], reverse=True)

    detector_by_attack = []
    for attack, rows in by_attack.items():
        detector_by_attack.append(
            {
                "attack": attack,
                "dr": mean([r["avg_dr"] for r in rows]),
                "fpr": mean([r["avg_fpr"] for r in rows]),
                "malicious_share": mean([r.get("avg_malicious_resource_share", float("nan")) for r in rows]),
            }
        )
    detector_by_attack.sort(key=lambda r: list(ATTACK_LABELS).index(r["attack"]))

    trust_rows = []
    for trust_mode, rows in group_by(summary_rows, "trust_mode").items():
        trust_rows.append(
            {
                "trust_mode": trust_mode,
                "mean_ratio": mean([r["avg_ratio"] for r in rows]),
                "mean_malicious_share": mean([r.get("avg_malicious_resource_share", float("nan")) for r in rows]),
                "mean_dr": mean([r["avg_dr"] for r in rows]),
                "mean_min_util": mean([r["avg_min_util"] for r in rows]),
            }
        )
    trust_rows.sort(key=lambda r: list(TRUST_MODE_LABELS).index(r["trust_mode"]))

    attack_trust_rows = []
    for attack in ATTACK_LABELS:
        for trust_mode in TRUST_MODE_LABELS:
            rows = [
                r for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ]
            if rows:
                attack_trust_rows.append(
                    {
                        "attack": attack,
                        "trust_mode": trust_mode,
                        "mean_ratio": mean([r["avg_ratio"] for r in rows]),
                        "mean_malicious_share": mean([r.get("avg_malicious_resource_share", float("nan")) for r in rows]),
                        "mean_dr": mean([r["avg_dr"] for r in rows]),
                    }
                )

    stealth_rows = []
    for attack in ["reputation_poison", "trust_mimicry"]:
        for trust_mode in TRUST_MODE_LABELS:
            rows = [
                r for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ]
            if rows:
                stealth_rows.append(
                    {
                        "attack": attack,
                        "trust_mode": trust_mode,
                        "mean_ratio": mean([r["avg_ratio"] for r in rows]),
                        "mean_malicious_share": mean([r.get("avg_malicious_resource_share", float("nan")) for r in rows]),
                        "mean_dr": mean([r["avg_dr"] for r in rows]),
                    }
                )

    by_comp_app = group_by(comparison_rows, "approach") if comparison_rows else {}
    horizon = []
    for app, rows in by_comp_app.items():
        horizon.append(
            {
                "approach": app,
                "mean_delta": mean([r["delta_400_minus_200"] for r in rows]),
                "burst_delta": next(
                    r["delta_400_minus_200"]
                    for r in rows
                    if r["attack_type"] == "burst"
                ),
                "stealth_delta": mean(
                    [
                        r["delta_400_minus_200"]
                        for r in rows
                        if r["attack_type"]
                        in {"reputation_poison", "trust_mimicry"}
                    ]
                ),
            }
        )
    horizon.sort(key=lambda r: r["mean_delta"], reverse=True)

    return {
        "app_rankings": app_rankings,
        "best_by_attack": best_by_attack,
        "lifts": lifts,
        "detector_by_attack": detector_by_attack,
        "trust_rows": trust_rows,
        "attack_trust_rows": attack_trust_rows,
        "stealth_rows": stealth_rows,
        "confidence_rows": paired_difference_rows(trial_rows or []),
        "horizon": horizon,
    }


def html_table(headers: list[str], rows: list[list[str]]) -> str:
    head = "".join(f"<th>{html.escape(h)}</th>" for h in headers)
    body_rows = []
    for row in rows:
        body_rows.append(
            "<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>"
        )
    return f"<table><thead><tr>{head}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"


def cite_html(ref_id: str) -> str:
    citation = next((c for c in CITATIONS if c["id"] == ref_id), None)
    if citation is None:
        return html.escape(ref_id)
    return f"<a href='#ref-{ref_id}' title='{html.escape(citation['label'])}'>[{ref_id}]</a>"


def source_map_html() -> str:
    rows = []
    for algorithm, origin, refs, note in SOURCE_MAP:
        ref_cells = []
        for ref in refs.split(" plus "):
            if ref in {c["id"] for c in CITATIONS}:
                ref_cells.append(cite_html(ref))
            else:
                ref_cells.append(html.escape(ref))
        rows.append(
            [
                html.escape(algorithm),
                html.escape(origin),
                " plus ".join(ref_cells),
                html.escape(note),
            ]
        )
    return html_table(["Approach / component", "Paper origin", "Citation", "Correct interpretation"], rows)


def references_html() -> str:
    items = []
    for citation in CITATIONS:
        items.append(
            f"<li id='ref-{citation['id']}'><strong>[{citation['id']}] "
            f"{html.escape(citation['label'])}.</strong> "
            f"{html.escape(citation['text'])} "
            f"<a href='{html.escape(citation['url'])}'>{html.escape(citation['url'])}</a></li>"
        )
    return "<ol class='refs'>" + "".join(items) + "</ol>"


def output_file_paths() -> list[Path]:
    interesting: list[Path] = []
    for base in [FULL_4500_DIR, ROOT / "output_400_rounds", ROOT / "outputs" / "Output_Moradi"]:
        if not base.exists():
            continue
        for path in sorted(base.rglob("*")):
            if path.is_file() and path.suffix.lower() in {".csv", ".txt", ".png"}:
                interesting.append(path)
    return interesting


def file_inventory() -> str:
    interesting = output_file_paths()
    items = "".join(
        f"<li><code>{html.escape(p.relative_to(ROOT).as_posix())}</code></li>"
        for p in interesting
    )
    return f"<details><summary>Output file inventory ({len(interesting)} files)</summary><ul>{items}</ul></details>"


def key_output_files() -> list[Path]:
    return [
        FULL_4500_CELL_CSV,
        FULL_4500_TRIAL_CSV,
        *[path for path, _ in PRIMARY_4500_FIGURES],
        SUMMARY_CSV,
        COMPARISON_CSV,
        RUN_NOTES,
        LEGACY_NOTES,
    ]


def reproducibility_files_html() -> str:
    rows = []
    for path in key_output_files():
        if path.exists():
            rows.append([
                html.escape(display_path(path)),
                html.escape(path.suffix.lower().lstrip(".") or "file"),
            ])
    return html_table(["File", "Type"], rows)


def graph_groups() -> list[tuple[str, list[Path]]]:
    groups: list[tuple[str, list[Path]]] = []
    legacy = sorted((ROOT / "outputs" / "Output_Moradi" / "outputs").glob("*.png"))
    if legacy:
        groups.append(("Original Moradi graph set", legacy))

    overview = sorted((ROOT / "output_400_rounds").glob("*.png"))
    if overview:
        groups.append(("400-round benchmark overview graphs", overview))

    for app in APPROACH_LABELS:
        folder = ROOT / "output_400_rounds" / app
        graphs = sorted(folder.glob("*.png"))
        if graphs:
            groups.append((APPROACH_LABELS[app], graphs))
    return groups


def all_graph_paths() -> list[Path]:
    paths: list[Path] = []
    for _, group in graph_groups():
        paths.extend(group)
    return paths


def is_valid_image(path: Path) -> bool:
    try:
        with Image.open(path) as img:
            img.verify()
        return True
    except Exception:
        return False


def text_size(draw: ImageDraw.ImageDraw, text: str, font) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def pil_save_bar(
    path: Path,
    title: str,
    labels: list[str],
    values: list[float],
    formatter,
    colors: list[str] | None = None,
) -> None:
    width, height = 1200, 720
    margin_l, margin_r, margin_t, margin_b = 100, 60, 90, 170
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.load_default(size=28)
    font = ImageFont.load_default(size=18)
    small = ImageFont.load_default(size=15)
    draw.text((margin_l, 28), title, fill="#18212f", font=title_font)
    chart_w = width - margin_l - margin_r
    chart_h = height - margin_t - margin_b
    y0 = margin_t + chart_h
    draw.line((margin_l, margin_t, margin_l, y0), fill="#44505c", width=2)
    draw.line((margin_l, y0, width - margin_r, y0), fill="#44505c", width=2)
    max_val = max(values) if values else 1.0
    max_val = max_val * 1.18 if max_val else 1.0
    bar_w = chart_w / max(1, len(values)) * 0.64
    colors = colors or ["#0b6b78"] * len(values)
    for i, value in enumerate(values):
        cx = margin_l + (i + 0.5) * chart_w / len(values)
        x1 = int(cx - bar_w / 2)
        x2 = int(cx + bar_w / 2)
        h = int(chart_h * value / max_val)
        y1 = y0 - h
        draw.rectangle((x1, y1, x2, y0), fill=colors[i], outline="#22313a")
        label = formatter(value)
        tw, _ = text_size(draw, label, small)
        draw.text((cx - tw / 2, y1 - 24), label, fill="#18212f", font=small)
        words = labels[i].replace(" hard gate", "").replace(" reference", "").split()
        for j, word in enumerate(words[:3]):
            tw, _ = text_size(draw, word, small)
            draw.text((cx - tw / 2, y0 + 14 + j * 18), word, fill="#354656", font=small)
    img.save(path)


def pil_save_horizontal_bar(path: Path, title: str, labels: list[str], ratios: list[float], shares: list[float]) -> None:
    width, height = 1300, 820
    margin_l, margin_r, margin_t, margin_b = 360, 80, 90, 60
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.load_default(size=28)
    font = ImageFont.load_default(size=16)
    small = ImageFont.load_default(size=15)
    draw.text((margin_l, 28), title, fill="#18212f", font=title_font)
    chart_w = width - margin_l - margin_r
    chart_h = height - margin_t - margin_b
    max_val = max(ratios) * 1.18 if ratios else 1
    gap = chart_h / max(1, len(ratios))
    bar_h = gap * 0.62
    for i, (label, ratio, share) in enumerate(zip(labels, ratios, shares)):
        cy = margin_t + (i + 0.5) * gap
        y1, y2 = int(cy - bar_h / 2), int(cy + bar_h / 2)
        x2 = int(margin_l + chart_w * ratio / max_val)
        draw.rectangle((margin_l, y1, x2, y2), fill="#0b6b78", outline="#22313a")
        draw.text((20, y1 + 2), label[:42], fill="#18212f", font=font)
        draw.text((x2 + 10, y1 + 1), f"{ratio:.3f} | mal {100*share:.1f}%", fill="#354656", font=small)
    draw.line((margin_l, margin_t, margin_l, height - margin_b), fill="#44505c", width=2)
    img.save(path)


def heat_color(value: float, lo: float = 0.25, hi: float = 0.85) -> tuple[int, int, int]:
    t = min(1.0, max(0.0, (value - lo) / (hi - lo)))
    start = (238, 245, 229)
    end = (11, 107, 120)
    return tuple(int(start[i] + t * (end[i] - start[i])) for i in range(3))


def pil_save_heatmap(path: Path, title: str, row_labels: list[str], col_labels: list[str], matrix: list[list[float]]) -> None:
    width, height = 1200, 780
    margin_l, margin_t = 220, 130
    cell_w, cell_h = 210, 92
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.load_default(size=28)
    font = ImageFont.load_default(size=17)
    small = ImageFont.load_default(size=15)
    draw.text((margin_l, 28), title, fill="#18212f", font=title_font)
    for j, label in enumerate(col_labels):
        x = margin_l + j * cell_w + 8
        draw.text((x, 86), label.replace(" hard gate", ""), fill="#18212f", font=small)
    for i, label in enumerate(row_labels):
        y = margin_t + i * cell_h + 34
        draw.text((24, y), label, fill="#18212f", font=font)
        for j, value in enumerate(matrix[i]):
            x1 = margin_l + j * cell_w
            y1 = margin_t + i * cell_h
            x2 = x1 + cell_w - 4
            y2 = y1 + cell_h - 4
            draw.rectangle((x1, y1, x2, y2), fill=heat_color(value), outline="white")
            txt = fmt_ratio(value)
            tw, th = text_size(draw, txt, font)
            draw.text((x1 + cell_w / 2 - tw / 2, y1 + cell_h / 2 - th / 2), txt, fill="#111820", font=font)
    img.save(path)


def pil_save_grouped_bar(
    path: Path,
    title: str,
    attacks: list[str],
    series: list[tuple[str, list[float], str]],
) -> None:
    width, height = 1300, 720
    margin_l, margin_r, margin_t, margin_b = 90, 60, 95, 150
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = ImageFont.load_default(size=28)
    font = ImageFont.load_default(size=16)
    small = ImageFont.load_default(size=14)
    draw.text((margin_l, 28), title, fill="#18212f", font=title_font)
    chart_w = width - margin_l - margin_r
    chart_h = height - margin_t - margin_b
    y0 = margin_t + chart_h
    draw.line((margin_l, margin_t, margin_l, y0), fill="#44505c", width=2)
    draw.line((margin_l, y0, width - margin_r, y0), fill="#44505c", width=2)
    group_w = chart_w / len(attacks)
    bar_w = group_w / (len(series) + 1)
    max_val = 0.9
    for s_idx, (name, vals, color) in enumerate(series):
        for i, value in enumerate(vals):
            x1 = margin_l + i * group_w + (s_idx + 0.45) * bar_w
            x2 = x1 + bar_w * 0.82
            y1 = y0 - chart_h * value / max_val
            draw.rectangle((x1, y1, x2, y0), fill=color, outline="#22313a")
    for i, attack in enumerate(attacks):
        x = margin_l + i * group_w + group_w / 2
        for j, word in enumerate(attack.split()):
            tw, _ = text_size(draw, word, small)
            draw.text((x - tw / 2, y0 + 12 + j * 18), word, fill="#354656", font=small)
    legend_x = margin_l + 600
    for i, (name, _, color) in enumerate(series):
        y = 38 + i * 26
        draw.rectangle((legend_x, y, legend_x + 20, y + 14), fill=color)
        draw.text((legend_x + 28, y - 2), name, fill="#18212f", font=font)
    img.save(path)


def generate_4500_figures_pil(summary_rows: list[dict], tables: dict) -> None:
    FULL_4500_FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    colors = {
        "algorithm1": "#0b6b78",
        "no_trust": "#9aa6b2",
        "oracle": "#2d7d46",
        "sparse_online": "#b45f06",
    }
    trust = tables["trust_rows"]
    pil_save_bar(
        PRIMARY_4500_FIGURES[0][0],
        "Trust-Mode Ablation: Welfare Ratio",
        [TRUST_MODE_LABELS[r["trust_mode"]] for r in trust],
        [r["mean_ratio"] for r in trust],
        fmt_ratio,
        [colors[r["trust_mode"]] for r in trust],
    )
    pil_save_bar(
        PRIMARY_4500_FIGURES[1][0],
        "Resource Leakage by Trust Mode",
        [TRUST_MODE_LABELS[r["trust_mode"]] for r in trust],
        [r["mean_malicious_share"] for r in trust],
        fmt_pct,
        [colors[r["trust_mode"]] for r in trust],
    )
    ranking = list(reversed(tables["app_rankings"]))
    pil_save_horizontal_bar(
        PRIMARY_4500_FIGURES[2][0],
        "Allocation Approach Ranking Under Algorithm 1",
        [APPROACH_LABELS[r["approach"]] for r in ranking],
        [r["mean_ratio"] for r in ranking],
        [r["mean_malicious_share"] for r in ranking],
    )
    attacks = list(ATTACK_LABELS)
    trust_modes = list(TRUST_MODE_LABELS)
    matrix = [
        [
            mean([
                r["avg_ratio"] for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ])
            for trust_mode in trust_modes
        ]
        for attack in attacks
    ]
    pil_save_heatmap(
        PRIMARY_4500_FIGURES[3][0],
        "Attack x Trust Mode: Mean NSW Ratio",
        [ATTACK_LABELS[a] for a in attacks],
        [TRUST_MODE_LABELS[t] for t in trust_modes],
        matrix,
    )
    selected = ["algorithm1", "sparse_online", "oracle"]
    series = []
    for trust_mode in selected:
        vals = [
            mean([
                r["avg_ratio"] for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ])
            for attack in attacks
        ]
        series.append((TRUST_MODE_LABELS[trust_mode], vals, colors[trust_mode]))
    pil_save_grouped_bar(
        PRIMARY_4500_FIGURES[4][0],
        "Algorithm 1 vs Sparse-Online vs Oracle Trust",
        [ATTACK_LABELS[a] for a in attacks],
        series,
    )


def generate_4500_figures(summary_rows: list[dict], tables: dict) -> None:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception as exc:
        print(f"Using PIL fallback for 4,500-trial figures: matplotlib unavailable ({exc})")
        generate_4500_figures_pil(summary_rows, tables)
        return

    FULL_4500_FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    colors = {
        "algorithm1": "#0b6b78",
        "no_trust": "#9aa6b2",
        "oracle": "#2d7d46",
        "sparse_online": "#b45f06",
    }

    def save(fig, path: Path) -> None:
        fig.tight_layout()
        fig.savefig(path, dpi=180, bbox_inches="tight")
        plt.close(fig)

    trust = tables["trust_rows"]
    labels = [TRUST_MODE_LABELS[r["trust_mode"]] for r in trust]
    x = list(range(len(trust)))

    fig, ax = plt.subplots(figsize=(8, 4.6))
    vals = [r["mean_ratio"] for r in trust]
    ax.bar(x, vals, color=[colors[r["trust_mode"]] for r in trust])
    ax.set_xticks(x, labels, rotation=18, ha="right")
    ax.set_ylabel("Mean NSW ratio")
    ax.set_title("Trust-Mode Ablation: Welfare Ratio")
    ax.set_ylim(0, max(vals) * 1.18)
    for i, val in enumerate(vals):
        ax.text(i, val + 0.012, fmt_ratio(val), ha="center", va="bottom", fontsize=9)
    save(fig, PRIMARY_4500_FIGURES[0][0])

    fig, ax = plt.subplots(figsize=(8, 4.6))
    vals = [r["mean_malicious_share"] for r in trust]
    ax.bar(x, vals, color=[colors[r["trust_mode"]] for r in trust])
    ax.set_xticks(x, labels, rotation=18, ha="right")
    ax.set_ylabel("Mean malicious resource share")
    ax.set_title("Resource Leakage by Trust Mode")
    ax.set_ylim(0, max(vals) * 1.22 if vals else 1)
    for i, val in enumerate(vals):
        ax.text(i, val + 0.006, fmt_pct(val), ha="center", va="bottom", fontsize=9)
    save(fig, PRIMARY_4500_FIGURES[1][0])

    ranking = list(reversed(tables["app_rankings"]))
    fig, ax = plt.subplots(figsize=(9, 5.8))
    y = list(range(len(ranking)))
    ratios = [r["mean_ratio"] for r in ranking]
    shares = [r["mean_malicious_share"] for r in ranking]
    ax.barh(y, ratios, color="#0b6b78")
    ax.set_yticks(y, [APPROACH_LABELS[r["approach"]] for r in ranking])
    ax.set_xlabel("Mean NSW ratio under Algorithm 1")
    ax.set_title("Allocation Approach Ranking Under Algorithm 1")
    ax.set_xlim(0, max(ratios) * 1.22)
    for i, (ratio, share) in enumerate(zip(ratios, shares)):
        ax.text(ratio + 0.01, i, f"{ratio:.3f} | mal {100*share:.1f}%", va="center", fontsize=8)
    save(fig, PRIMARY_4500_FIGURES[2][0])

    attacks = list(ATTACK_LABELS)
    trust_modes = list(TRUST_MODE_LABELS)
    matrix = []
    for attack in attacks:
        row = []
        for trust_mode in trust_modes:
            vals = [
                r["avg_ratio"] for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ]
            row.append(mean(vals))
        matrix.append(row)
    fig, ax = plt.subplots(figsize=(8.8, 5.4))
    im = ax.imshow(matrix, cmap="YlGnBu", vmin=0.25, vmax=0.85, aspect="auto")
    ax.set_xticks(range(len(trust_modes)), [TRUST_MODE_LABELS[t] for t in trust_modes], rotation=18, ha="right")
    ax.set_yticks(range(len(attacks)), [ATTACK_LABELS[a] for a in attacks])
    ax.set_title("Attack x Trust Mode: Mean NSW Ratio")
    for i, row in enumerate(matrix):
        for j, val in enumerate(row):
            ax.text(j, i, fmt_ratio(val), ha="center", va="center", fontsize=8, color="#12202b")
    fig.colorbar(im, ax=ax, label="Mean NSW ratio")
    save(fig, PRIMARY_4500_FIGURES[3][0])

    selected = ["algorithm1", "sparse_online", "oracle"]
    width = 0.24
    fig, ax = plt.subplots(figsize=(9, 5))
    base_x = list(range(len(attacks)))
    for offset, trust_mode in enumerate(selected):
        vals = []
        for attack in attacks:
            vals.append(mean([
                r["avg_ratio"] for r in summary_rows
                if r["attack_type"] == attack and r.get("trust_mode") == trust_mode
            ]))
        xs = [v + (offset - 1) * width for v in base_x]
        ax.bar(xs, vals, width=width, label=TRUST_MODE_LABELS[trust_mode], color=colors[trust_mode])
    ax.set_xticks(base_x, [ATTACK_LABELS[a] for a in attacks], rotation=18, ha="right")
    ax.set_ylabel("Mean NSW ratio")
    ax.set_title("Algorithm 1 vs Sparse-Online vs Oracle Trust")
    ax.legend(frameon=False)
    ax.set_ylim(0, 0.9)
    save(fig, PRIMARY_4500_FIGURES[4][0])


def graph_gallery_html() -> str:
    sections = []
    for title, paths in graph_groups():
        figures = []
        invalid = []
        for path in paths:
            caption = path.relative_to(ROOT).as_posix()
            if is_valid_image(path):
                figures.append(
                    "<figure class='thumb'>"
                    f"<a href='{rel(path)}'><img src='{rel(path)}' alt='{html.escape(caption)}'></a>"
                    f"<figcaption><code>{html.escape(caption)}</code></figcaption>"
                    "</figure>"
                )
            else:
                invalid.append(caption)
        invalid_note = ""
        if invalid:
            invalid_items = "".join(f"<li><code>{html.escape(path)}</code></li>" for path in invalid)
            invalid_note = (
                "<p class='warning'><strong>Excluded invalid/truncated image file(s):</strong></p>"
                f"<ul>{invalid_items}</ul>"
            )
        sections.append(
            f"<h3>{html.escape(title)} ({len(figures)} valid graphs"
            f"{', ' + str(len(invalid)) + ' excluded' if invalid else ''})</h3>"
            f"{invalid_note}<div class='gallery'>{''.join(figures)}</div>"
        )
    return "".join(sections)


def make_html(summary_rows: list[dict], comparison_rows: list[dict], tables: dict) -> str:
    generated = date.today().isoformat()
    n_trials = int(summary_rows[0]["n_trials_completed"])
    n_rounds = int(summary_rows[0]["n_rounds"])
    approach_count = len({r["approach"] for r in summary_rows})
    attack_count = len({r["attack_type"] for r in summary_rows})
    offline_nsw = mean([r["avg_offline_nsw"] for r in summary_rows])
    best = tables["app_rankings"][0]
    adaptive = next(r for r in tables["app_rankings"] if r["approach"] == "adaptive_alpha")
    baseline = next(r for r in tables["app_rankings"] if r["approach"] == "baseline_fixed_alpha")
    algorithm1_rows = [r for r in summary_rows if r.get("trust_mode") == "algorithm1"] or summary_rows

    ranking_rows = []
    for idx, row in enumerate(tables["app_rankings"], start=1):
        ranking_rows.append(
            [
                str(idx),
                html.escape(APPROACH_LABELS[row["approach"]]),
                fmt_ratio(row["mean_ratio"]),
                fmt_float(row["mean_nsw"]),
                fmt_float(row["mean_min_util"]),
                fmt_float(row["mean_fairness_gap"]),
                fmt_pct(row["mean_malicious_share"]),
            ]
        )

    non_extreme = [
        row for row in tables["app_rankings"]
        if row["approach"] not in {"pure_greedy_alpha0", "pure_equal_alpha1"}
    ]
    non_extreme_rows = [
        [
            str(idx),
            html.escape(APPROACH_LABELS[row["approach"]]),
            fmt_ratio(row["mean_ratio"]),
            fmt_float(row["mean_min_util"]),
            fmt_pct(row["mean_malicious_share"]),
        ]
        for idx, row in enumerate(non_extreme, start=1)
    ]

    attack_matrix = []
    by_attack = group_by(algorithm1_rows, "attack_type")
    for attack in ATTACK_LABELS:
        rows = sorted(by_attack[attack], key=lambda r: r["avg_ratio"], reverse=True)
        best_app = rows[0]["approach"]
        cells = [html.escape(ATTACK_LABELS[attack])]
        for app in APPROACH_LABELS:
            val = next(r for r in rows if r["approach"] == app)["avg_ratio"]
            text = fmt_ratio(val)
            if app == best_app:
                text = f"<strong>{text}</strong>"
            cells.append(text)
        attack_matrix.append(cells)

    lift_rows = [
        [
            html.escape(APPROACH_LABELS[r["approach"]]),
            fmt_float(r["abs_lift"]),
            fmt_pct(r["rel_lift"]),
        ]
        for r in tables["lifts"]
    ]

    detector_rows = [
        [
            html.escape(ATTACK_LABELS[r["attack"]]),
            fmt_pct(r["dr"]),
            fmt_pct(r["fpr"]),
            fmt_pct(r["malicious_share"]),
        ]
        for r in tables["detector_by_attack"]
    ]

    trust_rows = [
        [
            html.escape(TRUST_MODE_LABELS[r["trust_mode"]]),
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
            fmt_float(r["mean_min_util"]),
        ]
        for r in tables["trust_rows"]
    ]

    attack_trust_rows = [
        [
            html.escape(ATTACK_LABELS[r["attack"]]),
            html.escape(TRUST_MODE_LABELS[r["trust_mode"]]),
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
        ]
        for r in tables["attack_trust_rows"]
    ]

    stealth_rows = [
        [
            html.escape(ATTACK_LABELS[r["attack"]]),
            html.escape(TRUST_MODE_LABELS[r["trust_mode"]]),
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
        ]
        for r in tables["stealth_rows"]
    ]

    horizon_rows = [
        [
            html.escape(APPROACH_LABELS[r["approach"]]),
            fmt_float(r["mean_delta"]),
            fmt_float(r["burst_delta"]),
            fmt_float(r["stealth_delta"]),
        ]
        for r in tables["horizon"]
    ]

    confidence_rows = [
        [
            html.escape(r["comparison"]),
            fmt_float(r["mean_diff"]),
            f"+/- {fmt_float(r['ci95'])}",
            str(r["n_pairs"]),
        ]
        for r in tables["confidence_rows"]
    ]

    approach_cards = []
    for item in APPROACH_GUIDE:
        key = item["key"]
        approach_cards.append(
            f"<section class='approach'><h3>{html.escape(APPROACH_LABELS[key])}</h3>"
            f"<p class='role'><strong>{html.escape(item['role'])}.</strong> "
            f"{html.escape(APPROACH_DESCRIPTIONS[key])}</p>"
            f"<p><strong>Intuition and implementation.</strong> {html.escape(item['intuition'])} "
            f"{html.escape(item['implementation'])}</p>"
            f"<p><strong>Interpretation and caveat.</strong> {html.escape(item['interpretation'])}</p>"
            f"</section>"
        )

    figure_html = []
    for path, caption in PRIMARY_4500_FIGURES:
        if path.exists():
            figure_html.append(
                f"<figure><img src='{rel(path)}' alt='{html.escape(caption)}'>"
                f"<figcaption>{html.escape(caption)}</figcaption></figure>"
            )

    legacy_summary = html.escape(LEGACY_NOTES.read_text(errors="replace") if LEGACY_NOTES.exists() else "")
    run_notes = html.escape(RUN_NOTES.read_text(errors="replace") if RUN_NOTES.exists() else "")

    css = """
    :root { color-scheme: light; --ink:#18212f; --muted:#607086; --line:#d9e0ea; --accent:#0b6b78; --soft:#eef7f6; }
    body { margin:0; font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif; color:var(--ink); background:#fbfcfe; line-height:1.55; }
    header { padding:56px 7vw 40px; background:#0d2730; color:#fff; }
    header p { max-width:980px; color:#d5e7ea; font-size:18px; }
    main { max-width:1120px; margin:0 auto; padding:34px 24px 80px; }
    h1 { margin:0 0 16px; font-size:40px; line-height:1.1; letter-spacing:0; }
    h2 { margin:42px 0 12px; font-size:25px; border-bottom:1px solid var(--line); padding-bottom:8px; }
    h3 { margin:18px 0 8px; font-size:18px; }
    .meta { display:flex; gap:12px; flex-wrap:wrap; margin-top:20px; }
    .pill { border:1px solid rgba(255,255,255,.35); padding:6px 10px; border-radius:999px; color:#eaf4f6; font-size:14px; }
    .callout { background:var(--soft); border-left:4px solid var(--accent); padding:16px 18px; margin:22px 0; }
    table { width:100%; border-collapse:collapse; margin:16px 0 24px; font-size:14px; background:white; }
    th, td { border:1px solid var(--line); padding:8px 10px; vertical-align:top; }
    th { background:#eef2f7; text-align:left; }
    code { background:#eef2f7; padding:1px 4px; border-radius:4px; }
    figure { margin:24px 0 34px; background:white; border:1px solid var(--line); padding:12px; }
    img { width:100%; height:auto; display:block; }
    figcaption { color:var(--muted); font-size:13px; margin-top:8px; }
    .grid { display:grid; grid-template-columns:repeat(auto-fit,minmax(260px,1fr)); gap:12px; }
    .gallery { display:grid; grid-template-columns:repeat(auto-fit,minmax(300px,1fr)); gap:14px; }
    .thumb { margin:0; padding:10px; }
    .thumb img { aspect-ratio: 4 / 3; object-fit:contain; background:#fff; }
    .missing-box { aspect-ratio:4 / 3; display:flex; align-items:center; justify-content:center; background:#f5ecec; color:#8a2424; text-align:center; padding:12px; }
    .approach { background:white; border:1px solid var(--line); padding:14px 16px; }
    .approach p { color:#39485c; }
    .approach .role { color:#24364a; }
    .refs li { margin:0 0 10px; }
    .warning { color:#8a2424; }
    pre { white-space:pre-wrap; background:#101820; color:#e9f4f6; padding:14px; overflow:auto; font-size:12px; }
    details { margin:18px 0; }
    summary { cursor:pointer; font-weight:600; }
    """

    return f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Trust-Aware Online NSW Allocation Manuscript</title>
  <style>{css}</style>
</head>
<body>
<header>
  <h1>Trust-Aware Online Nash Welfare Allocation Under Strategic Attacks</h1>
    <p>A manuscript-style explanation of the Moradi final simulator, the allocator variants, and the completed 4,500-trial trust-ablation benchmark.</p>
  <div class="meta">
    <span class="pill">{generated}</span>
    <span class="pill">{approach_count} allocation variants</span>
    <span class="pill">{attack_count} attack types</span>
    <span class="pill">{n_trials} trials per cell</span>
    <span class="pill">{n_rounds} rounds</span>
  </div>
</header>
<main>
  <section>
    <h2>Abstract</h2>
    <p>This report studies how allocation rules from online welfare-maximization literature behave when they are placed behind an imperfect trust layer. The Moradi final simulator combines trust detection {cite_html("R2")} with a trust-filtered set-aside greedy allocator based on online Nash social welfare with predictions {cite_html("R1")}. The completed benchmark evaluates {approach_count} simulator-level allocation variants, {attack_count} attack types, and four trust modes: no trust, Algorithm 1, oracle trust, and sparse-online trust. The evaluated variants are not full reproductions of all original source-paper methods. The reported welfare ratio uses a Frank-Wolfe numerical approximation {cite_html("R7")} for the same concave NSW objective. This makes the denominator faster and more stable than relying on an SLSQP solve that may fall back to equal split, but it is still a numerical approximation rather than a formal certificate of optimality.</p>
    <div class="callout"><strong>Main result.</strong> Under the Algorithm 1 hard trust gate, {html.escape(APPROACH_LABELS[best["approach"]])} achieved the highest empirical NSW ratio among the tested approach variants, with mean ratio {fmt_ratio(best["mean_ratio"])}. This is an efficiency reference, not the proposed robust method: it also sends {fmt_pct(best["mean_malicious_share"])} of allocation to malicious agents on average. Among non-extreme project methods, {html.escape(APPROACH_LABELS["expert_advice"])} is the strongest empirical controller and {html.escape(APPROACH_LABELS["adaptive_alpha"])} is the cleanest direct modification, with mean ratio {fmt_ratio(adaptive["mean_ratio"])} versus {fmt_ratio(baseline["mean_ratio"])} for the fixed-alpha baseline.</div>
    <div class="callout"><strong>Main takeaway.</strong> Allocation changes help, but trust quality is still the dominant bottleneck under stealth attacks. Expert advice gives the strongest empirical controller, while adaptive alpha is the cleanest direct extension of the original fixed-alpha allocator. Adaptive alpha does not win every setting; its value is that it changes the baseline allocation rule directly rather than wrapping several allocators or relying on a stronger trust gate.</div>
  </section>

  <section>
    <h2>Overview</h2>
    <p>The simulator studies online divisible-good allocation when some agents strategically misreport their values. Each round has one unit of resource; the algorithm observes reported values, applies a trust gate, and allocates only through the information available at that time. Evaluation uses true values, so malicious reports can help attackers only if they survive the trust layer and influence allocation.</p>
    <p>The main metric is Nash social welfare (NSW), the geometric mean of legitimate agents' cumulative utilities, following the online NSW-with-predictions setting of Banerjee et al. {cite_html("R1")}. The reported NSW ratio compares the online policy against a clairvoyant offline benchmark; higher ratios mean the online allocation is closer to the offline benchmark.</p>
  </section>

  <section>
    <h2>Project Vocabulary</h2>
    {html_table(["Term", "Meaning"], [
      ["Agent", "A participant who may receive part of the resource each round."],
      ["Legitimate agent", "An honest agent. These are the agents whose welfare matters in the final NSW score."],
      ["Malicious agent", "An attacker that may lie in reports and try to manipulate trust."],
      ["Report", "The value announced to the allocator. It may be false for malicious agents."],
      ["True valuation", "The value used to compute real utility after the allocation is made."],
      ["Trusted set", "Agents currently allowed to receive resource from the allocator."],
      ["Detection rate", "How often malicious agents are identified."],
      ["False positive rate", "How often legitimate agents are incorrectly flagged."],
      ["Offline benchmark", "A stronger comparison policy that sees the whole sequence before allocating."],
    ])}
  </section>

  <section>
    <h2>Problem Formulation</h2>
    <p>The simulator can be summarized as an online divisible-good allocation problem with strategic reports and trust filtering. The report uses this notation throughout.</p>
    {html_table(["Object", "Formula / definition"], [[html.escape(a), html.escape(b)] for a, b in PROBLEM_FORMULATION])}
  </section>

  <section>
    <h2>Round-by-Round Logic</h2>
    <p>Each simulation round follows the same basic pipeline. First, attackers decide whether to attack and create distorted reports. Second, the trust detector updates cumulative trust evidence and forms a trusted set. Third, the allocator divides the unit budget among trusted agents. Fourth, the simulator updates utilities using true valuations rather than reports. Finally, the code logs welfare, detection rate, false positives, minimum utility, fairness gap, malicious resource share, and other trust/allocation diagnostics. The older Moradi-focused runs also include detector F1-style summaries, while the combined 4,500-trial benchmark emphasizes detection rates and resource leakage.</p>
    <p>This structure is important because a good detector can protect allocation, but only if the malicious agents are removed before allocation. Stealth attacks are dangerous because they can survive long enough to affect future allocation state.</p>
  </section>

  <section>
    <h2>Research Thesis</h2>
    <p>The main goal of this report is to study how different allocation variants behave under a shared trust-aware simulator. We use Algorithm 1 as the main trust setting, then include no-trust, oracle-trust, and sparse-online modes as ablations.</p>
    <p>This means the report is not mainly trying to propose a new trust detector. The trust layer defines the experimental environment: sometimes the allocator receives clean input, and sometimes it receives contaminated input from trusted-but-malicious agents. The research question is which allocation rule is most robust under those conditions.</p>
  </section>

  <section>
    <h2>Why Compare These Variants Together?</h2>
    <p>These methods do not all solve the exact same theoretical problem in their original papers. We compare them because they can all be expressed as allocation rules or allocation controllers inside the same online trust-filtered simulator: each method receives the same reports, the same trusted set, and the same per-round budget, then outputs an allocation. This creates a fair empirical harness for studying how allocation behavior changes when the trust layer is imperfect.</p>
  </section>

  <section>
    <h2>Original Project Directions and Scope</h2>
    <p>The original project had three natural directions. This report focuses mainly on the third one: allocation alternatives under a shared trust-aware setting.</p>
    {html_table(["Direction", "Question", "Role in this report"], [[html.escape(c) for c in row] for row in ORIGINAL_PROJECT_DIRECTIONS])}
  </section>

  <section>
    <h2>Experimental Design</h2>
    <p>This compact design block is the main experiment in one place. It is the bridge between the paper ideas and the empirical benchmark.</p>
    {html_table(["Item", "Setting"], [
      ["Agents", "10 legitimate agents and 4 malicious agents"],
      ["Rounds", f"{n_rounds} online allocation rounds per trial"],
      ["Trials", f"{n_trials} trials per cell"],
      ["Allocation variants", f"{approach_count} variants, including two alpha-reference endpoints"],
      ["Attack types", f"{attack_count} attack types"],
      ["Trust modes", "4 trust modes: no trust, Algorithm 1, oracle trust, sparse-online"],
      ["Total cells", f"{approach_count} x {attack_count} x 4 = {approach_count * attack_count * 4} cells"],
      ["Total trials", f"{approach_count * attack_count * 4} x {n_trials} = {approach_count * attack_count * 4 * n_trials} trials"],
      ["Offline denominator", "Frank-Wolfe numerical approximation to offline NSW"],
      ["Main metrics", "NSW ratio, malicious resource share, detection rate, minimum utility"],
    ])}
  </section>

  <section>
    <h2>Trust Modes</h2>
    <p>The trust modes are not separate allocation algorithms. They define how much contamination reaches the allocator, which is why they are essential for interpreting allocation performance.</p>
    {html_table(["Trust mode", "Meaning", "Purpose"], [[html.escape(c) for c in row] for row in TRUST_MODE_OVERVIEW])}
  </section>

  <section>
    <h2>Attack Types</h2>
    {html_table(["Attack", "Plain-language explanation"], [
      ["Byzantine", "A direct and noisy manipulation pattern. It is easy for the detector to catch."],
      ["Burst", "An attacker behaves normally most of the time, then attacks in bursts. Detection eventually catches it, but recovery takes longer."],
      ["Reputation poison", "An attacker behaves well early to build trust credit, then attacks later. This exploits cumulative trust inertia."],
      ["Trust mimicry", "An attacker tries to stay close enough to normal behavior to avoid detection while still damaging welfare."],
      ["Compound", "A coordinated attack type present in the code. It was not part of the 400-round approach benchmark, but it is included in the combined 4,500-trial benchmark."],
    ])}
  </section>

  <section>
    <h2>What Was Evaluated</h2>
    <p>The primary dataset is the combined full 4,500-trial summary: 9 allocation variants x 5 attacks x 4 trust modes x {n_trials} trials. Each trial uses {n_rounds} rounds and logs welfare ratio, true legitimate NSW, minimum utility, fairness gap, detection rates, and malicious resource share, with valuation seeds held fixed across variants.</p>
    <p>The 400-round all-approaches folder and Moradi focused output folder are treated as appendix context. They are useful for auditability and visual diagnostics, while the manuscript tables and claims below use the combined 4,500-trial benchmark as the primary evidence base.</p>
    <div class="callout"><strong>Interpretation boundary.</strong> We compare seven trust-aware allocation variants plus two alpha-reference endpoints under the same experimental harness, not nine methods with identical theoretical assumptions. Although these variants originate from different theoretical frameworks, we compare them inside the same simulation environment. Therefore, the results should be interpreted as empirical performance under a shared benchmark, not as a claim that one paper's theorem dominates another paper's theorem.</div>
  </section>

  <section>
    <h2>Methods</h2>
    <p>The simulator allocates divisible budget each round under strategic reporting. For this report, the existing trust mechanism is used as the experimental environment. The allocation variants are evaluated after this trust filter, so detector results are included mainly to explain when the allocator receives clean input versus contaminated input. The new variant files share the same simulation step through TrustModeStepMixin, so attack timing, detection updates, metric logging, and offline comparison are held consistent across variants. The beta-gap trust detector is based on Akgun, Aydin, Gil, and Nedic {cite_html("R2")}, while the Moradi graph diagnostics adapt trust-based sparsest-subgraph clustering {cite_html("R3")}.</p>
    <p>The evaluated variant set is fixed-alpha baseline, adaptive alpha, PACE-inspired trusted allocation, generalized-mean greedy, sample-resolving approximation, expert advice, robust aggregation, and two alpha-reference endpoints: pure greedy alpha=0 and pure equal alpha=1.</p>
  </section>

  <section>
    <h2>How To Read The Metrics</h2>
    {html_table(["Metric", "How to interpret it"], [
      ["Average NSW", "The achieved geometric-mean utility for legitimate agents. Higher is better."],
      ["NSW ratio", "Achieved NSW divided by offline benchmark NSW. Higher is better; 1.0 would match the offline clairvoyant benchmark."],
      ["Minimum utility", "The worst cumulative utility among legitimate agents. Higher means the most harmed honest agent is better protected."],
      ["Fairness gap", "Spread between better-off and worse-off legitimate agents. Lower is usually more equal, but can also reflect everyone receiving less."],
      ["Detection rate", "Fraction of malicious agents detected. Higher is better."],
      ["False positive rate", "Fraction of legitimate agents wrongly detected. Lower is better."],
      ["Convergence round", "Approximate round when detection stabilizes. Smaller is faster."],
      ["Alg1 / Spectral / Sparse F1", "Different detector-quality summaries. In this report, sparse F1 is diagnostic unless used online by the allocator."],
    ])}
  </section>

  <section>
    <h2>Important Clarification</h2>
    <div class="callout"><strong>Important clarification.</strong> The allocation variants below are not all full reproductions of their source papers. The goal is not to prove that one paper's original theorem dominates another. Instead, we implement a set of comparable allocation variants inside the same trust-aware simulator. Some variants are direct adaptations, some are paper-inspired approximations, and some are project-defined baselines or controllers. This makes the comparison empirically useful, but the claims should be interpreted as simulator-level evidence rather than theorem-level comparison across papers.</div>
  </section>

  <section>
    <h2>From Source Papers to Simulator Variants</h2>
    <p>This is the literature-to-code translation layer. The source papers provide allocation ideas under their own assumptions; the simulator asks how those ideas behave after a trust gate, strategic reports, and shared NSW-ratio evaluation are imposed.</p>
    {html_table(["Variant", "Original paper idea", "Original assumptions", "Implemented here", "Interpretation"], [[html.escape(c) for c in row] for row in SOURCE_TO_SIMULATOR_MAPPING])}
  </section>

  <section>
    <h2>Allocation Variants</h2>
    <p>The explanations below slow down the comparison before the numerical tables. Each variant is described in terms of its intuition, how it is implemented in this simulator, and how its result should be interpreted. This matters because the rows are not all the same kind of object: some are allocators, some are controllers, some are approximations, and two are reference endpoints.</p>
    <div class="grid">{''.join(approach_cards)}</div>
  </section>

  <section>
    <h2>Allocation Variant Taxonomy</h2>
    <p>This shorter taxonomy is separate from the citation map. It makes clear that the rows are not all the same theoretical object: some are allocators, some are controllers, one is a meta-controller, and one is a preprocessing heuristic.</p>
    {html_table(["Approach", "Type", "Full paper reproduction?", "Used as"], [[html.escape(c) for c in row] for row in APPROACH_TAXONOMY])}
  </section>

  <section>
    <h2>Main Contribution vs Baseline</h2>
    <p>The main project contribution is not that every new variant beats the literature baseline. The clean contribution is the adaptive-alpha controller: it keeps the set-aside allocator recognizable while using trust instability to control how defensive the allocation should be.</p>
    {html_table(["Component", "Existing baseline", "Project contribution"], [[html.escape(c) for c in row] for row in CONTRIBUTION_TABLE])}
  </section>

  <section>
    <h2>Results</h2>
    <p><strong>How to read the results.</strong> The results are organized in three steps. First, compare allocation variants under the Algorithm 1 hard trust gate. Second, remove the alpha endpoints to focus on non-extreme methods rather than pure efficiency or pure defense references. Third, use trust-mode ablations to separate allocator effects from trust-gate effects.</p>
    <p><strong>Efficiency references versus robustness contributions.</strong> Pure greedy is included to show the high-efficiency endpoint of the alpha tradeoff, but because it has no set-aside protection, it is not treated as the proposed robust method.</p>
    <h3>Ranking Under Algorithm 1 Hard Trust Gate</h3>
    {html_table(["Rank", "Approach", "Mean ratio", "Mean NSW", "Mean min utility", "Mean fairness gap", "Mean malicious share"], ranking_rows)}
    <p><strong>Ranking caveat.</strong> This ranking is empirical and implementation-specific. The alpha=0 pure greedy row is a reference endpoint, not the proposed robust method. It achieves high NSW by emphasizing immediate efficiency, but it also has the largest mean malicious resource share in the Algorithm 1 ranking. That makes it an upper-efficiency anchor, not the final recommendation.</p>
    <h3>Algorithm 1 Ranking, Non-Extreme Methods Only</h3>
    <p>This table removes the alpha=0 and alpha=1 reference endpoints so the main allocation-method comparison is not visually dominated by pure greedy.</p>
    {html_table(["Rank", "Approach", "Mean ratio", "Mean min utility", "Mean malicious share"], non_extreme_rows)}
    <h3>Ratio by Attack Under Algorithm 1</h3>
    {html_table(["Attack"] + [html.escape(APPROACH_LABELS[a]) for a in APPROACH_LABELS], attack_matrix)}
    <h3>Lift Over Fixed-Alpha Baseline Under Algorithm 1</h3>
    <p>This lift is computed only within the Algorithm 1 hard-gate condition, using the fixed-alpha set-aside row for the same attack as the denominator. It is not averaged across no-trust, oracle, or sparse-online trust modes.</p>
    {html_table(["Approach", "Mean absolute ratio lift", "Mean relative lift"], lift_rows)}
    <h3>Trust-Mode Ablation</h3>
    {html_table(["Trust mode", "Mean ratio", "Mean malicious share", "Detection rate", "Mean min utility"], trust_rows)}
    <h3>Attack by Trust Mode</h3>
    {html_table(["Attack", "Trust mode", "Mean ratio", "Mean malicious share", "Detection rate"], attack_trust_rows)}
    <h3>Stealth-Attack Resource Leakage</h3>
    {html_table(["Attack", "Trust mode", "Mean ratio", "Mean malicious share", "Detection rate"], stealth_rows)}
    <h3>Paired 95% Confidence Intervals</h3>
    <p>These intervals use the per-trial CSV, pairing runs by attack and trial id; the sparse-online comparison is also paired by approach. They are descriptive uncertainty intervals for this simulator, not formal theorem-level guarantees.</p>
    {html_table(["Comparison", "Mean ratio difference", "95% CI half-width", "Paired observations"], confidence_rows)}
  </section>

  <section>
    <h2>Key Findings</h2>
    {html_table(["Finding", "Meaning"], [[html.escape(c) for c in row] for row in KEY_FINDINGS])}
  </section>

  <section>
    <h2>Why Expert Advice Is the Strongest Non-Extreme Controller</h2>
    <p>Expert advice is the strongest non-extreme controller here because it is not a standalone allocation rule; it is a controller that selects among allocation behaviors implemented in this simulator. It uses several candidate allocators and shifts weight toward the candidate that performs better under the current attack pattern. In this implementation, those candidates include defensive set-aside, aggressive set-aside, and generalized-mean behavior.</p>
    <p>Fixed-alpha set-aside commits to one alpha value for the entire run. That makes it clean and interpretable, but it cannot automatically become more aggressive during easier periods or more defensive when the attack pattern changes. The expert-advice controller can adapt across those behaviors, which explains why it can outperform the single fixed-alpha rule in this mixed-attack simulator without implying that the expert-advice paper itself solves the original online NSW allocation problem.</p>
  </section>

  <section>
    <h2>Why Adaptive Alpha Is The Main Contribution</h2>
    <p>Expert advice is strongest empirically, but it is harder to claim as a clean direct contribution because it is a meta-controller over several candidate allocators. Adaptive alpha is more defensible as the main proposed method because it minimally extends the original set-aside allocator and uses the trust layer to control robustness. Its empirical advantage should be read as trust-mode dependent rather than universal.</p>
    <p>This distinction also makes the paper narrative cleaner: fixed-alpha set-aside is the directly sourced baseline, adaptive alpha is the proposed trust-aware extension, and expert advice is an empirical controller benchmark showing how much can be gained by adaptively combining several behaviors.</p>
  </section>

  <section>
    <h2>Mechanism Evidence</h2>
    <p>The trust-ablation results directly separate detector effects from allocator effects. No-trust has mean ratio {fmt_ratio(next(r["mean_ratio"] for r in tables["trust_rows"] if r["trust_mode"] == "no_trust"))} and malicious resource share {fmt_pct(next(r["mean_malicious_share"] for r in tables["trust_rows"] if r["trust_mode"] == "no_trust"))}. Algorithm 1 improves the mean ratio to {fmt_ratio(next(r["mean_ratio"] for r in tables["trust_rows"] if r["trust_mode"] == "algorithm1"))}, but still leaks {fmt_pct(next(r["mean_malicious_share"] for r in tables["trust_rows"] if r["trust_mode"] == "algorithm1"))} of allocation to malicious agents on average. Oracle trust reaches {fmt_ratio(next(r["mean_ratio"] for r in tables["trust_rows"] if r["trust_mode"] == "oracle"))}; sparse-online trust nearly matches it at {fmt_ratio(next(r["mean_ratio"] for r in tables["trust_rows"] if r["trust_mode"] == "sparse_online"))}. This is a trust-gate finding, not the main allocation contribution.</p>
    <p>The sparse-online result is promising, but it may benefit from simulator structure. In this implementation, the sparsest-subgraph detector is given <code>cfg.n_malicious</code> and trims the detected set to that known malicious count, which is a strong cardinality prior. Therefore, sparse-online should be interpreted as a diagnostic trust-gate upper-ablation between Algorithm 1 and oracle trust, not as a deployment-ready detector. It should not be claimed as theoretically guaranteed, and it needs sensitivity tests under noisier trust observations, different malicious-agent counts, and less separable attack behavior before being presented as a final trust mechanism.</p>
    <p>The mechanism is clearest on Reputation Poison and Trust Mimicry: Algorithm 1 has zero final detection rate on both stealth attacks, and its welfare is essentially the same as no trust. Sparse-online and oracle trust remove nearly all malicious resource share and restore ratios near the oracle benchmark. This means the allocator comparison and the trust mechanism must be interpreted together.</p>
  </section>

  <section>
    <h2>Detector Behavior</h2>
    <p>Algorithm 1 from the trustworthy-consensus line {cite_html("R2")} catches obvious Byzantine and Burst behavior under the hard trust gate, partially catches Compound, and misses the two stealth attacks. The sparse-online gate adapts the Moradi trust-graph/sparsest-subgraph idea {cite_html("R3")} as an online allocation gate; in this experiment it behaves almost like oracle trust, but it should still be described as an empirical simulator result rather than a deployment-ready theorem.</p>
    <p>Allocation consequence: when Algorithm 1 fails on stealth attacks, hard-gate trust stops helping. Welfare then depends on whether the allocator can limit damage from agents that are still labeled trusted.</p>
    {html_table(["Attack", "Algorithm 1 detection rate", "False positive rate", "Mean malicious share"], detector_rows)}
  </section>

  <section>
    <h2>Horizon Effects</h2>
    <p>This horizon-effect table comes from the earlier 200-vs-400 comparison output, not from the combined 4,500-trial grid, which uses 400 rounds throughout.</p>
    <p>Moving from 200 to 400 rounds helps Burst most because delayed detection and recovery have more time to amortize. Stealth attacks barely improve for static or sample-heavy variants, showing that longer horizons alone do not solve reputation poisoning.</p>
    {html_table(["Approach", "Mean delta", "Burst delta", "Stealth delta"], horizon_rows)}
  </section>

  <section>
    <h2>Interpretation</h2>
    <p>The benchmark supports two conclusions. First, a stronger offline benchmark lowers ratios relative to an equal-split fallback, but it is the right comparison because it makes the denominator a real offline optimum approximation rather than an accidental weak baseline. Second, robustness should be added at the controller level before claiming that a full allocator replacement is superior. Expert advice performs well because it is a meta-controller: it adaptively combines several allocator candidates rather than committing to one fixed allocation rule. In this implementation, those candidates include set-aside-style experts and generalized-mean behavior, and the weights are updated using the exponential-weights idea {cite_html("R6")}.</p>
    <p>PACE-inspired {cite_html("R4")}, generalized-mean {cite_html("R5")}, robust-aggregation {cite_html("R9")}, and sample-resolving {cite_html("R5")} variants should be treated as implementation-level baselines or approximations, not full reproductions of the corresponding theoretical methods. Under Algorithm 1, none of them beat the fixed-alpha baseline on mean ratio in this run. The negative results for PACE-inspired, generalized-mean, and sample-resolving variants should be read as results for these simplified implementations, not as failures of the original papers.</p>
  </section>

  <section>
    <h2>Conclusion</h2>
    <p>Under Algorithm 1, pure greedy is the highest-efficiency endpoint but not the safest robustness contribution because it relies heavily on the trust gate and leaks the most resource to malicious agents. If the trust gate is wrong, pure greedy has no set-aside cushion, so the same aggressiveness that raises NSW can also route resource to attackers. Expert advice is the strongest practical controller because it can mix allocator behaviors. The direct alpha modification remains useful because it changes the fixed-alpha set-aside baseline itself, although its advantage depends on trust quality. PACE-inspired, generalized-mean, sample-resolving, and robust-aggregation variants do not improve over fixed-alpha in this simulator. When trust improves to sparse-online or oracle trust, most allocation methods improve, showing that trust quality and allocation rule interact strongly.</p>
    <p>These findings should be interpreted as evidence about the implemented simulator variants, not as a theoretical ranking of the original source papers.</p>
  </section>

  <section>
    <h2>Final Recommendation</h2>
    {html_table(["Role", "Choice", "Reason"], [[html.escape(c) for c in row] for row in FINAL_RECOMMENDATION_TABLE])}
  </section>

  <section>
    <h2>Primary 4,500-Trial Visualizations</h2>
    <p>The five figures below are created directly from <code>full_4500_cell_summary.csv</code> and embedded in the main report body so that the visual evidence matches the primary benchmark.</p>
    {''.join(figure_html)}
  </section>

  <section>
    <h2>Reproducibility Files</h2>
    <p>The full legacy graph folders are kept as reproducibility artifacts rather than embedded into the main report. This compact table lists the files needed to reproduce the main tables, figures, and appendix context.</p>
    {reproducibility_files_html()}
  </section>

  <section>
    <h2>Appendix: Allocation Variant Comparison</h2>
    <p>This reference table is kept outside the main narrative because it overlaps with the variant explanations and taxonomy. It is useful for auditability: it states what each variant changes relative to the fixed-alpha set-aside baseline and why it might succeed or fail after the trust filter.</p>
    {html_table(["Approach", "Is it a true allocator?", "How it uses trust", "What changes from baseline", "Expected strength", "Expected weakness"], [[html.escape(c) for c in row] for row in ALLOCATION_ALGORITHM_COMPARISON])}
  </section>

  <section>
    <h2>Appendix: Citation Map</h2>
    <p>This attribution table lists the paper sources and project-defined components. It is kept in the appendix because the main body now explains the paper-to-simulator translation directly.</p>
    {source_map_html()}
  </section>

  <section>
    <h2>Limitations and Threats to Validity</h2>
    <ul>
      <li>The report summarizes completed CSV outputs; it is not a formal proof of correctness.</li>
      <li>The Frank-Wolfe offline benchmark is faster and more stable than relying on an SLSQP equal-split fallback, but it is still a numerical approximation rather than a formal optimality certificate.</li>
      <li>The trust observation process uses simulator-side status and attack information, so detector performance should not be presented as deployment-ready evidence.</li>
      <li>The sparse-online trust result is strong in this simulator, but it should be treated as an empirical gate comparison rather than a theoretical guarantee.</li>
      <li>GPU acceleration is not expected to materially help this code, because the workload is Python, NumPy, SciPy, and plotting/document generation CPU work rather than tensor computation.</li>
    </ul>
  </section>

  <section>
    <h2>References</h2>
    {references_html()}
  </section>

  <section>
    <h2>Appendix: Output Files Used</h2>
    <p>Final interpretation: fixed-alpha set-aside greedy is the cleanest directly sourced theoretical baseline; adaptive alpha is the cleanest direct allocation-side extension; expert advice is the strongest empirical controller; and pure greedy alpha=0 should be treated as an efficiency reference rather than the proposed robust method.</p>
    <ul>
      <li><code>{html.escape(display_path(FULL_4500_CELL_CSV))}</code></li>
      <li><code>{html.escape(display_path(FULL_4500_TRIAL_CSV))}</code></li>
      <li><code>{html.escape(str(SUMMARY_CSV.relative_to(ROOT)))}</code></li>
      <li><code>{html.escape(str(COMPARISON_CSV.relative_to(ROOT)))}</code></li>
      <li><code>{html.escape(str(RUN_NOTES.relative_to(ROOT)))}</code></li>
      <li><code>{html.escape(str(LEGACY_NOTES.relative_to(ROOT)))}</code></li>
    </ul>
  </section>
</main>
</body>
</html>
"""


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def add_picture_if_exists(document: Document, path: Path, caption: str) -> None:
    if not path.exists():
        return
    if not is_valid_image(path):
        document.add_paragraph(f"{caption} (image file exists but appears truncated, so it was not embedded)")
        return
    document.add_picture(str(path), width=Inches(6.4))
    p = document.add_paragraph(caption)
    p.style = "Caption"


def make_docx(summary_rows: list[dict], comparison_rows: list[dict], tables: dict) -> None:
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    document.add_heading("Trust-Aware Online Nash Welfare Allocation Under Strategic Attacks", 0)
    document.add_paragraph(
        "Manuscript/report based on the Moradi final simulator outputs and "
        "the combined full 4,500-trial trust-ablation benchmark."
    )
    document.add_paragraph(date.today().isoformat())

    n_trials = int(summary_rows[0]["n_trials_completed"])
    n_rounds = int(summary_rows[0]["n_rounds"])
    best = tables["app_rankings"][0]
    adaptive = next(r for r in tables["app_rankings"] if r["approach"] == "adaptive_alpha")
    baseline = next(r for r in tables["app_rankings"] if r["approach"] == "baseline_fixed_alpha")

    document.add_heading("Abstract", 1)
    document.add_paragraph(
        "This report studies how allocation rules from online welfare-maximization "
        "literature behave when they are placed behind an imperfect trust layer. "
        f"The benchmark evaluates {len(APPROACH_LABELS)} simulator-level allocation variants across "
        f"{len(ATTACK_LABELS)} attack types and four trust modes, with {n_trials} "
        f"completed trials per cell and {n_rounds} rounds per trial. "
        "The evaluated variants are not full reproductions of all original source-paper methods. "
        "The reported welfare ratio uses a Frank-Wolfe numerical approximation for the "
        "same concave NSW objective. This makes the denominator faster and more stable "
        "than relying on an SLSQP solve that may fall back to equal split, but it is still "
        "a numerical approximation rather than a formal certificate of optimality [R7]."
    )
    document.add_paragraph(
        f"Under the Algorithm 1 hard trust gate, {APPROACH_LABELS[best['approach']]} "
        f"achieved the highest empirical NSW ratio, with mean ratio "
        f"{fmt_ratio(best['mean_ratio'])}. This is an efficiency reference, not the "
        f"proposed robust method: it also sends {fmt_pct(best['mean_malicious_share'])} "
        f"of allocation to malicious agents on average. Among non-extreme project methods, "
        f"expert advice is the strongest empirical controller and adaptive alpha is the "
        f"cleanest direct modification, with mean ratio {fmt_ratio(adaptive['mean_ratio'])} "
        f"versus {fmt_ratio(baseline['mean_ratio'])} for the fixed-alpha baseline."
    )
    document.add_paragraph(
        "Main takeaway: allocation changes help, but trust quality is still the "
        "dominant bottleneck under stealth attacks. Expert advice gives the strongest "
        "empirical controller, while adaptive alpha is the cleanest direct extension "
        "of the original fixed-alpha allocator. Adaptive alpha does not win every "
        "setting; its value is that it changes the baseline allocation rule directly "
        "rather than wrapping several allocators or relying on a stronger trust gate."
    )

    document.add_heading("Overview", 1)
    document.add_paragraph(
        "The simulator studies online divisible-good allocation when some agents "
        "strategically misreport their values. Each round has one unit of resource; "
        "the algorithm observes reported values, applies a trust gate, and allocates "
        "only through information available at that time. Evaluation uses true values, "
        "so malicious reports can help attackers only if they survive the trust layer "
        "and influence allocation."
    )
    document.add_paragraph(
        "The main metric is Nash social welfare, or NSW, the geometric mean of "
        "legitimate agents' cumulative utilities. The NSW ratio compares the online "
        "policy against a clairvoyant offline benchmark. The set-aside plus greedy "
        "prediction structure follows the online NSW-with-predictions framework of "
        "Banerjee et al. [R1]."
    )
    document.add_heading("Project Vocabulary", 1)
    add_table(
        document,
        ["Term", "Meaning"],
        [
            ["Agent", "A participant who may receive part of the resource each round."],
            ["Legitimate agent", "An honest agent. These agents define the welfare score."],
            ["Malicious agent", "An attacker that can lie in reports or manipulate trust."],
            ["Report", "The value announced to the allocation algorithm."],
            ["True valuation", "The value used after allocation to compute real utility."],
            ["Trusted set", "Agents currently allowed to receive resource."],
            ["Offline benchmark", "A clairvoyant comparison policy used as the ratio denominator."],
        ],
    )

    document.add_heading("Problem Formulation", 1)
    document.add_paragraph(
        "The simulator can be written as an online divisible-good allocation problem "
        "with strategic reports and trust filtering. This notation is the mathematical "
        "core of the experiment."
    )
    add_table(
        document,
        ["Object", "Formula / definition"],
        PROBLEM_FORMULATION,
    )

    document.add_heading("Round-by-Round Logic", 1)
    steps = [
        "Attackers decide whether to attack and may distort their reports.",
        "The trust detector updates cumulative trust evidence.",
        "The simulator builds the trusted set from the detector outputs.",
        "The allocator divides the unit budget among trusted agents.",
        "Utilities are updated using true valuations, not reported valuations.",
        "The code logs welfare, detection rate, false positives, minimum utility, fairness gap, malicious resource share, and other trust/allocation diagnostics.",
    ]
    for step in steps:
        document.add_paragraph(step, style="List Number")
    document.add_paragraph(
        "The older Moradi-focused runs also include detector F1-style summaries, while "
        "the combined 4,500-trial benchmark emphasizes detection rates and resource leakage."
    )

    document.add_heading("Research Thesis", 1)
    document.add_paragraph(
        "The main goal of this report is to study how different allocation variants "
        "behave under a shared trust-aware simulator. We use Algorithm 1 as the main "
        "trust setting, then include no-trust, oracle-trust, and sparse-online modes "
        "as ablations."
    )
    document.add_paragraph(
        "This means the report is not mainly trying to propose a new trust detector. "
        "The trust layer defines the experimental environment: sometimes the allocator "
        "receives clean input, and sometimes it receives contaminated input from "
        "trusted-but-malicious agents. The research question is which allocation rule "
        "is most robust under those conditions."
    )

    document.add_heading("Why Compare These Variants Together?", 1)
    document.add_paragraph(
        "These methods do not all solve the exact same theoretical problem in their "
        "original papers. We compare them because they can all be expressed as "
        "allocation rules or allocation controllers inside the same online trust-filtered "
        "simulator: each method receives the same reports, the same trusted set, and "
        "the same per-round budget, then outputs an allocation. This creates a fair "
        "empirical harness for studying how allocation behavior changes when the trust "
        "layer is imperfect."
    )

    document.add_heading("Original Project Directions and Scope", 1)
    document.add_paragraph(
        "The original project had three natural directions. This report focuses mainly "
        "on the third one: allocation alternatives under a shared trust-aware setting."
    )
    add_table(
        document,
        ["Direction", "Question", "Role in this report"],
        ORIGINAL_PROJECT_DIRECTIONS,
    )

    document.add_heading("Experimental Design", 1)
    document.add_paragraph(
        "This compact design block is the main experiment in one place. It is the bridge "
        "between the paper ideas and the empirical benchmark."
    )
    add_table(
        document,
        ["Item", "Setting"],
        [
            ["Agents", "10 legitimate agents and 4 malicious agents"],
            ["Rounds", f"{n_rounds} online allocation rounds per trial"],
            ["Trials", f"{n_trials} trials per cell"],
            ["Allocation variants", f"{len(APPROACH_LABELS)} variants, including two alpha-reference endpoints"],
            ["Attack types", f"{len(ATTACK_LABELS)} attack types"],
            ["Trust modes", "4 trust modes: no trust, Algorithm 1, oracle trust, sparse-online"],
            ["Total cells", f"{len(APPROACH_LABELS)} x {len(ATTACK_LABELS)} x 4 = {len(APPROACH_LABELS) * len(ATTACK_LABELS) * 4} cells"],
            ["Total trials", f"{len(APPROACH_LABELS) * len(ATTACK_LABELS) * 4} x {n_trials} = {len(APPROACH_LABELS) * len(ATTACK_LABELS) * 4 * n_trials} trials"],
            ["Offline denominator", "Frank-Wolfe numerical approximation to offline NSW"],
            ["Main metrics", "NSW ratio, malicious resource share, detection rate, minimum utility"],
        ],
    )

    document.add_heading("Trust Modes", 1)
    document.add_paragraph(
        "The trust modes are not separate allocation algorithms. They define how much "
        "contamination reaches the allocator, which is why they are essential for "
        "interpreting allocation performance."
    )
    add_table(document, ["Trust mode", "Meaning", "Purpose"], TRUST_MODE_OVERVIEW)

    document.add_heading("Attack Types", 1)
    add_table(
        document,
        ["Attack", "Plain-language explanation"],
        [
            ["Byzantine", "Direct noisy manipulation. It is the easiest attack for the detector to catch."],
            ["Burst", "The attacker behaves normally, then attacks in concentrated bursts."],
            ["Reputation Poison", "The attacker behaves well early, builds trust credit, then attacks later."],
            ["Trust Mimicry", "The attacker stays close enough to normal behavior to avoid detection while still hurting welfare."],
            ["Compound", "A coordinated attack type in the Moradi final code. It was not part of the 400-round all-approach benchmark, but it is included in the combined 4,500-trial benchmark."],
        ],
    )

    document.add_heading("What Was Evaluated", 1)
    document.add_paragraph(
        f"The primary dataset is the combined full 4,500-trial summary: 9 allocation variants "
        f"times 5 attack types times 4 trust modes times {n_trials} trials. Each trial "
        f"uses {n_rounds} rounds and logs welfare ratio, true legitimate NSW, minimum "
        "utility, fairness gap, detection rates, and malicious resource share, with "
        "valuation seeds held fixed across variants."
    )
    document.add_paragraph(
        "The 400-round all-approaches folder and Moradi focused output folder are "
        "treated as appendix context. They are useful for auditability and visual "
        "diagnostics, while the manuscript tables and claims use the combined "
        "4,500-trial benchmark as the evidence base."
    )
    document.add_paragraph(
        "Interpretation boundary: we compare seven trust-aware allocation variants "
        "plus two alpha-reference endpoints under the same experimental harness, not "
        "nine methods with identical theoretical assumptions. Although these variants "
        "originate from different theoretical frameworks, the results should be interpreted as empirical "
        "performance under a shared benchmark, not as a claim that one paper's theorem "
        "dominates another paper's theorem."
    )

    document.add_heading("Methods", 1)
    document.add_paragraph(
        "The simulator allocates divisible budget each round under strategic reporting. "
        "For this report, the existing trust mechanism is used as the experimental "
        "environment. The allocation variants are evaluated after this trust filter, "
        "so detector results are included mainly to explain when the allocator receives "
        "clean input versus contaminated input. The new approach files share the same "
        "simulation step through TrustModeStepMixin, so attack timing, detection updates, "
        "metric logging, and offline comparison are held consistent across variants. "
        "The beta-gap trust detector is based on Akgun, Aydin, Gil, and Nedic [R2], "
        "while the Moradi graph diagnostics adapt trust-based sparsest-subgraph "
        "clustering [R3]."
    )
    document.add_paragraph(
        "The attack set is Byzantine, Burst, Reputation Poison, Trust Mimicry, and Compound. "
        "The evaluated variant set is fixed-alpha baseline, adaptive alpha, PACE-inspired "
        "trusted allocation, generalized-mean greedy, sample-resolving approximation, "
        "expert advice, robust aggregation, and two alpha-reference endpoints."
    )
    document.add_heading("How To Read The Metrics", 1)
    add_table(
        document,
        ["Metric", "How to interpret it"],
        [
            ["Average NSW", "Achieved geometric-mean utility for legitimate agents. Higher is better."],
            ["NSW ratio", "Achieved NSW divided by the offline benchmark NSW. Higher is better."],
            ["Minimum utility", "Worst cumulative utility among legitimate agents. Higher means better protection for the most harmed honest agent."],
            ["Fairness gap", "Spread between better-off and worse-off legitimate agents. Lower is more equal, but can also mean everyone is similarly low."],
            ["Detection rate", "How often malicious agents are detected. Higher is better."],
            ["False positive rate", "How often legitimate agents are wrongly detected. Lower is better."],
            ["Malicious resource share", "Average fraction of allocation sent to malicious agents. Lower means less resource leakage."],
        ],
    )

    document.add_heading("Important Clarification", 1)
    document.add_paragraph(
        "Important clarification: the allocation variants below are not all full reproductions "
        "of their source papers. The goal is not to prove that one paper's original "
        "theorem dominates another. Instead, we implement a set of comparable allocation "
        "variants inside the same trust-aware simulator. Some variants are direct "
        "adaptations, some are paper-inspired approximations, and some are project-defined "
        "baselines or controllers. This makes the comparison empirically useful, but "
        "the claims should be interpreted as simulator-level evidence rather than "
        "theorem-level comparison across papers."
    )

    document.add_heading("From Source Papers to Simulator Variants", 1)
    document.add_paragraph(
        "This is the literature-to-code translation layer. The source papers provide "
        "allocation ideas under their own assumptions; the simulator asks how those "
        "ideas behave after a trust gate, strategic reports, and shared NSW-ratio "
        "evaluation are imposed."
    )
    add_table(
        document,
        ["Variant", "Original paper idea", "Original assumptions", "Implemented here", "Interpretation"],
        SOURCE_TO_SIMULATOR_MAPPING,
    )

    document.add_heading("Allocation Variants", 1)
    document.add_paragraph(
        "The explanations below slow down the comparison before the numerical tables. "
        "Each variant is described in terms of its intuition, how it is implemented "
        "in this simulator, and how its result should be interpreted. This matters "
        "because the rows are not all the same kind of object: some are allocators, "
        "some are controllers, some are approximations, and two are reference endpoints."
    )
    for item in APPROACH_GUIDE:
        key = item["key"]
        document.add_heading(APPROACH_LABELS[key], 2)
        document.add_paragraph(f"{item['role']}. {APPROACH_DESCRIPTIONS[key]}")
        document.add_paragraph(f"Intuition and implementation: {item['intuition']} {item['implementation']}")
        document.add_paragraph(f"Interpretation and caveat: {item['interpretation']}")

    document.add_heading("Allocation Variant Taxonomy", 1)
    document.add_paragraph(
        "This shorter taxonomy is separate from the citation map. It makes clear that "
        "the rows are not all the same theoretical object: some are allocators, some "
        "are controllers, one is a meta-controller, and one is a preprocessing heuristic."
    )
    add_table(
        document,
        ["Approach", "Type", "Full paper reproduction?", "Used as"],
        APPROACH_TAXONOMY,
    )

    document.add_heading("Main Contribution vs Baseline", 1)
    document.add_paragraph(
        "The main project contribution is not that every new variant beats the "
        "literature baseline. The clean contribution is the adaptive-alpha controller: "
        "it keeps the set-aside allocator recognizable while using trust instability "
        "to control how defensive the allocation should be."
    )
    add_table(
        document,
        ["Component", "Existing baseline", "Project contribution"],
        CONTRIBUTION_TABLE,
    )

    document.add_heading("Results", 1)
    document.add_paragraph(
        "How to read the results: the results are organized in three steps. First, "
        "compare allocation variants under the Algorithm 1 hard trust gate. Second, "
        "remove the alpha endpoints to focus on non-extreme methods rather than pure "
        "efficiency or pure defense references. Third, use trust-mode ablations to "
        "separate allocator effects from trust-gate effects."
    )
    document.add_paragraph(
        "Efficiency references versus robustness contributions: pure greedy is included "
        "to show the high-efficiency endpoint of the alpha tradeoff, but because it has "
        "no set-aside protection, it is not treated as the proposed robust method."
    )
    document.add_heading("Ranking Under Algorithm 1 Hard Trust Gate", 2)
    ranking_rows = [
        [
            str(idx),
            APPROACH_LABELS[r["approach"]],
            fmt_ratio(r["mean_ratio"]),
            fmt_float(r["mean_nsw"]),
            fmt_float(r["mean_min_util"]),
            fmt_float(r["mean_fairness_gap"]),
            fmt_pct(r["mean_malicious_share"]),
        ]
        for idx, r in enumerate(tables["app_rankings"], start=1)
    ]
    add_table(
        document,
        ["Rank", "Approach", "Mean ratio", "Mean NSW", "Mean min utility", "Mean fairness gap", "Mean malicious share"],
        ranking_rows,
    )
    document.add_paragraph(
        "Ranking caveat: this ranking is empirical and implementation-specific. The "
        "alpha=0 pure greedy row is a reference endpoint, not the proposed robust method. "
        "It achieves high NSW by emphasizing immediate efficiency, but it also has the "
        "largest mean malicious resource share in the Algorithm 1 ranking. That makes it "
        "an upper-efficiency anchor, not the final recommendation."
    )
    document.add_heading("Algorithm 1 Ranking, Non-Extreme Methods Only", 2)
    non_extreme = [
        row for row in tables["app_rankings"]
        if row["approach"] not in {"pure_greedy_alpha0", "pure_equal_alpha1"}
    ]
    non_extreme_rows = [
        [
            str(idx),
            APPROACH_LABELS[row["approach"]],
            fmt_ratio(row["mean_ratio"]),
            fmt_float(row["mean_min_util"]),
            fmt_pct(row["mean_malicious_share"]),
        ]
        for idx, row in enumerate(non_extreme, start=1)
    ]
    add_table(
        document,
        ["Rank", "Approach", "Mean ratio", "Mean min utility", "Mean malicious share"],
        non_extreme_rows,
    )

    document.add_heading("Ratio by Attack Under Algorithm 1", 2)
    algorithm1_rows = [r for r in summary_rows if r.get("trust_mode") == "algorithm1"] or summary_rows
    by_attack = group_by(algorithm1_rows, "attack_type")
    attack_rows = []
    for attack in ATTACK_LABELS:
        rows = sorted(by_attack[attack], key=lambda r: r["avg_ratio"], reverse=True)
        attack_rows.append(
            [
                ATTACK_LABELS[attack],
                APPROACH_LABELS[rows[0]["approach"]],
                fmt_ratio(rows[0]["avg_ratio"]),
                fmt_ratio(rows[1]["avg_ratio"]),
                fmt_ratio(rows[-1]["avg_ratio"]),
            ]
        )
    add_table(
        document,
        ["Attack", "Best approach", "Best ratio", "Second-best ratio", "Worst ratio"],
        attack_rows,
    )

    document.add_heading("Lift Over Fixed-Alpha Baseline Under Algorithm 1", 2)
    document.add_paragraph(
        "This lift is computed only within the Algorithm 1 hard-gate condition, using "
        "the fixed-alpha set-aside row for the same attack as the denominator. It is not "
        "averaged across no-trust, oracle, or sparse-online trust modes."
    )
    lift_rows = [
        [APPROACH_LABELS[r["approach"]], fmt_float(r["abs_lift"]), fmt_pct(r["rel_lift"])]
        for r in tables["lifts"]
    ]
    add_table(document, ["Approach", "Mean absolute ratio lift", "Mean relative lift"], lift_rows)

    document.add_heading("Trust-Mode Ablation", 2)
    trust_rows = [
        [
            TRUST_MODE_LABELS[r["trust_mode"]],
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
            fmt_float(r["mean_min_util"]),
        ]
        for r in tables["trust_rows"]
    ]
    add_table(
        document,
        ["Trust mode", "Mean ratio", "Mean malicious share", "Detection rate", "Mean min utility"],
        trust_rows,
    )

    document.add_heading("Attack by Trust Mode", 2)
    attack_trust_rows = [
        [
            ATTACK_LABELS[r["attack"]],
            TRUST_MODE_LABELS[r["trust_mode"]],
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
        ]
        for r in tables["attack_trust_rows"]
    ]
    add_table(
        document,
        ["Attack", "Trust mode", "Mean ratio", "Mean malicious share", "Detection rate"],
        attack_trust_rows,
    )

    document.add_heading("Stealth-Attack Resource Leakage", 2)
    stealth_rows = [
        [
            ATTACK_LABELS[r["attack"]],
            TRUST_MODE_LABELS[r["trust_mode"]],
            fmt_ratio(r["mean_ratio"]),
            fmt_pct(r["mean_malicious_share"]),
            fmt_pct(r["mean_dr"]),
        ]
        for r in tables["stealth_rows"]
    ]
    add_table(
        document,
        ["Attack", "Trust mode", "Mean ratio", "Mean malicious share", "Detection rate"],
        stealth_rows,
    )

    document.add_heading("Paired 95% Confidence Intervals", 2)
    document.add_paragraph(
        "These intervals use the per-trial CSV, pairing runs by attack and trial id; "
        "the sparse-online comparison is also paired by approach. They are descriptive "
        "uncertainty intervals for this simulator, not formal theorem-level guarantees."
    )
    confidence_rows = [
        [
            r["comparison"],
            fmt_float(r["mean_diff"]),
            f"+/- {fmt_float(r['ci95'])}",
            str(r["n_pairs"]),
        ]
        for r in tables["confidence_rows"]
    ]
    add_table(
        document,
        ["Comparison", "Mean ratio difference", "95% CI half-width", "Paired observations"],
        confidence_rows,
    )

    document.add_heading("Key Findings", 1)
    add_table(document, ["Finding", "Meaning"], KEY_FINDINGS)

    document.add_heading("Why Expert Advice Is the Strongest Non-Extreme Controller", 1)
    document.add_paragraph(
        "Expert advice is the strongest non-extreme controller here because it is not "
        "a standalone allocation rule; it is a controller that selects among allocation "
        "behaviors implemented in this simulator. "
        "It uses several candidate allocators and shifts weight toward the "
        "candidate that performs better under the current attack pattern. In this "
        "implementation, those candidates include defensive set-aside, aggressive "
        "set-aside, and generalized-mean behavior."
    )
    document.add_paragraph(
        "Fixed-alpha set-aside commits to one alpha value for the entire run. That makes "
        "it clean and interpretable, but it cannot automatically become more aggressive "
        "during easier periods or more defensive when the attack pattern changes. The "
        "expert-advice controller can adapt across those behaviors, which explains why "
        "it can outperform the single fixed-alpha rule in this mixed-attack simulator "
        "without implying that the expert-advice paper itself solves the original online "
        "NSW allocation problem."
    )

    document.add_heading("Why Adaptive Alpha Is The Main Contribution", 1)
    document.add_paragraph(
        "Expert advice is strongest empirically, but it is harder to claim as a clean "
        "direct contribution because it is a meta-controller over several candidate "
        "allocators. Adaptive alpha is more defensible as the main proposed method "
        "because it minimally extends the original set-aside allocator and uses the "
        "trust layer to control robustness. Its empirical advantage should be read as "
        "trust-mode dependent rather than universal."
    )
    document.add_paragraph(
        "This distinction also makes the paper narrative cleaner: fixed-alpha set-aside "
        "is the directly sourced baseline, adaptive alpha is the proposed trust-aware "
        "extension, and expert advice is an empirical controller benchmark showing how "
        "much can be gained by adaptively combining several behaviors."
    )

    document.add_heading("Mechanism Evidence", 1)
    document.add_paragraph(
        f"The trust-ablation results separate detector effects from allocator effects. "
        f"No-trust has mean ratio {fmt_ratio(next(r['mean_ratio'] for r in tables['trust_rows'] if r['trust_mode'] == 'no_trust'))} "
        f"and malicious resource share {fmt_pct(next(r['mean_malicious_share'] for r in tables['trust_rows'] if r['trust_mode'] == 'no_trust'))}. "
        f"Algorithm 1 improves the mean ratio to {fmt_ratio(next(r['mean_ratio'] for r in tables['trust_rows'] if r['trust_mode'] == 'algorithm1'))}, "
        f"but still leaks {fmt_pct(next(r['mean_malicious_share'] for r in tables['trust_rows'] if r['trust_mode'] == 'algorithm1'))} "
        f"of allocation to malicious agents on average. Oracle trust reaches "
        f"{fmt_ratio(next(r['mean_ratio'] for r in tables['trust_rows'] if r['trust_mode'] == 'oracle'))}; "
        f"sparse-online trust nearly matches it at "
        f"{fmt_ratio(next(r['mean_ratio'] for r in tables['trust_rows'] if r['trust_mode'] == 'sparse_online'))}. "
        "This is a trust-gate finding, not the main allocation contribution."
    )
    document.add_paragraph(
        "The sparse-online result is promising, but it may benefit from simulator "
        "structure. In this implementation, the sparsest-subgraph detector is given "
        "cfg.n_malicious and trims the detected set to that known malicious count, which "
        "is a strong cardinality prior. Therefore, sparse-online should be interpreted "
        "as a diagnostic trust-gate upper-ablation between Algorithm 1 and oracle trust, "
        "not as a deployment-ready detector. It should not be claimed as theoretically "
        "guaranteed, and it needs sensitivity tests under noisier trust observations, "
        "different malicious-agent counts, and less separable attack behavior before "
        "being presented as a final trust mechanism."
    )
    document.add_paragraph(
        "The mechanism is clearest on Reputation Poison and Trust Mimicry: Algorithm 1 "
        "has zero final detection rate on both stealth attacks, and its welfare is "
        "essentially the same as no trust. Sparse-online and oracle trust remove nearly "
        "all malicious resource share and restore ratios near the oracle benchmark. This "
        "means the allocator comparison and the trust mechanism must be interpreted together."
    )

    document.add_heading("Detector Behavior", 1)
    document.add_paragraph(
        "Algorithm 1 catches the obvious Byzantine and Burst attacks under the hard "
        "trust gate, partially catches Compound, and misses the two stealth attacks. The "
        "sparse-online gate adapts the Moradi trust-graph/sparsest-subgraph idea [R3] "
        "as an online allocation gate; in this experiment it behaves almost like oracle "
        "trust, but it should still be described as an empirical simulator result rather "
        "than a deployment-ready theorem."
    )
    document.add_paragraph(
        "Allocation consequence: when Algorithm 1 fails on stealth attacks, hard-gate "
        "trust stops helping. Welfare then depends on whether the allocator can limit "
        "damage from agents that are still labeled trusted."
    )
    detector_rows = [
        [
            ATTACK_LABELS[r["attack"]],
            fmt_pct(r["dr"]),
            fmt_pct(r["fpr"]),
            fmt_pct(r["malicious_share"]),
        ]
        for r in tables["detector_by_attack"]
    ]
    add_table(
        document,
        ["Attack", "Algorithm 1 detection rate", "FPR", "Mean malicious share"],
        detector_rows,
    )

    document.add_heading("Horizon Effects", 1)
    document.add_paragraph(
        "This horizon-effect table comes from the earlier 200-vs-400 comparison output, "
        "not from the combined 4,500-trial grid, which uses 400 rounds throughout."
    )
    document.add_paragraph(
        "The 200-vs-400 comparison shows that Burst benefits most from a longer horizon, "
        "while stealth attacks mostly plateau unless the approach has an adaptive or "
        "expert-control mechanism."
    )
    horizon_rows = [
        [
            APPROACH_LABELS[r["approach"]],
            fmt_float(r["mean_delta"]),
            fmt_float(r["burst_delta"]),
            fmt_float(r["stealth_delta"]),
        ]
        for r in tables["horizon"]
    ]
    add_table(document, ["Approach", "Mean delta", "Burst delta", "Stealth delta"], horizon_rows)

    document.add_heading("Interpretation", 1)
    document.add_paragraph(
        "The benchmark supports using expert advice as the empirical controller benchmark "
        "among implemented variants and the adaptive alpha rule as the most direct "
        "extension of the fixed-alpha baseline. The offline benchmark is better because it finds a stronger "
        "offline allocation; lower ratios are therefore more honest, not worse evidence. "
        "Expert advice should be interpreted as a meta-controller, not as a direct "
        "replacement allocator from 1997. It adaptively combines several allocator "
        "candidates, including set-aside-style experts, and uses the exponential-weights "
        "idea of Freund and Schapire [R6]."
    )
    document.add_paragraph(
        "PACE-inspired, generalized-mean, robust-aggregation, and sample-resolving "
        "variants remain useful comparators. However, under Algorithm 1 they do "
        "not beat the fixed-alpha baseline on average ratio. These variants should be "
        "treated as implementation-level baselines or approximations, not full "
        "reproductions of the corresponding theoretical methods. The negative results "
        "for PACE-inspired, generalized-mean, and sample-resolving variants should be "
        "read as results for these simplified implementations, not as failures of the "
        "original papers."
    )

    document.add_heading("Conclusion", 1)
    document.add_paragraph(
        "Under Algorithm 1, pure greedy is the highest-efficiency endpoint but not the "
        "safest robustness contribution because it relies heavily on the trust gate and "
        "leaks the most resource to malicious agents. If the trust gate is wrong, pure "
        "greedy has no set-aside cushion, so the same aggressiveness that raises NSW can "
        "also route resource to attackers. Expert advice is the strongest practical "
        "controller because it can mix allocator behaviors. The direct alpha modification "
        "remains useful because it changes the fixed-alpha set-aside baseline itself, "
        "although its advantage depends on trust quality. PACE-inspired, generalized-mean, "
        "sample-resolving, and robust-aggregation variants do not improve over fixed-alpha "
        "in this simulator. When trust improves to sparse-online or oracle trust, most "
        "allocation methods improve, showing that trust quality and allocation rule "
        "interact strongly."
    )
    document.add_paragraph(
        "These findings should be interpreted as evidence about the implemented "
        "simulator variants, not as a theoretical ranking of the original source papers."
    )

    document.add_heading("Final Recommendation", 1)
    add_table(document, ["Role", "Choice", "Reason"], FINAL_RECOMMENDATION_TABLE)

    document.add_heading("Primary 4,500-Trial Visualizations", 1)
    document.add_paragraph(
        "The five figures below are created directly from full_4500_cell_summary.csv "
        "and embedded in the main report body so that the visual evidence matches the "
        "primary benchmark."
    )
    for path, caption in PRIMARY_4500_FIGURES:
        add_picture_if_exists(document, path, caption)

    document.add_heading("Reproducibility Files", 1)
    document.add_paragraph(
        "The full legacy graph folders are kept as reproducibility artifacts rather "
        "than embedded into the main report. This compact table lists the files needed "
        "to reproduce the main tables, figures, and appendix context."
    )
    key_rows = [
        [display_path(path), path.suffix.lower().lstrip(".") or "file"]
        for path in key_output_files()
        if path.exists()
    ]
    add_table(document, ["File", "Type"], key_rows)

    document.add_heading("Appendix: Allocation Variant Comparison", 1)
    document.add_paragraph(
        "This reference table is kept outside the main narrative because it overlaps "
        "with the variant explanations and taxonomy. It is useful for auditability: "
        "it states what each variant changes relative to the fixed-alpha set-aside "
        "baseline and why it might succeed or fail after the trust filter."
    )
    add_table(
        document,
        [
            "Approach",
            "Is it a true allocator?",
            "How it uses trust",
            "What changes from baseline",
            "Expected strength",
            "Expected weakness",
        ],
        ALLOCATION_ALGORITHM_COMPARISON,
    )

    document.add_heading("Appendix: Citation Map", 1)
    document.add_paragraph(
        "This attribution table lists the paper sources and project-defined components. "
        "It is kept in the appendix because the main body now explains the paper-to-simulator "
        "translation directly."
    )
    add_table(
        document,
        ["Approach / component", "Paper origin", "Citation", "Correct interpretation"],
        SOURCE_MAP,
    )

    document.add_heading("Limitations and Threats to Validity", 1)
    limitations = [
        "This report summarizes completed CSV outputs; it is not a formal proof of correctness.",
        "The Frank-Wolfe offline benchmark is faster and more stable than relying on an SLSQP equal-split fallback, but it is still a numerical approximation rather than a formal optimality certificate.",
        "The trust observation process is simulator-assisted and should not be described as deployment-ready detection evidence.",
        "The sparse-online trust result is strong in this simulator, but it should be treated as an empirical gate comparison rather than a theoretical guarantee.",
        "GPU acceleration is unlikely to help much because the workload is CPU-bound Python, NumPy, SciPy, and plotting/document generation.",
    ]
    for item in limitations:
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("References", 1)
    for citation in CITATIONS:
        document.add_paragraph(f"[{citation['id']}] {citation['text']} {citation['url']}")

    document.add_heading("Appendix: Output Files Used", 1)
    document.add_paragraph(
        "Final interpretation: fixed-alpha set-aside greedy is the cleanest directly "
        "sourced theoretical baseline; adaptive alpha is the cleanest direct allocation-side "
        "extension; expert advice is the strongest empirical controller; and pure greedy "
        "alpha=0 should be treated as an efficiency reference rather than the proposed "
        "robust method."
    )
    if FULL_4500_CELL_CSV.exists():
        document.add_paragraph(display_path(FULL_4500_CELL_CSV))
    if FULL_4500_TRIAL_CSV.exists():
        document.add_paragraph(display_path(FULL_4500_TRIAL_CSV))
    document.add_paragraph(str(SUMMARY_CSV.relative_to(ROOT)))
    document.add_paragraph(str(COMPARISON_CSV.relative_to(ROOT)))
    document.add_paragraph(str(RUN_NOTES.relative_to(ROOT)))
    document.add_paragraph(str(LEGACY_NOTES.relative_to(ROOT)))

    DOCS.mkdir(exist_ok=True)
    document.save(DOCX_OUT)


def main() -> None:
    if FULL_4500_CELL_CSV.exists():
        summary_rows = normalize_full_4500_rows(read_csv(FULL_4500_CELL_CSV))
    else:
        summary_rows = read_csv(SUMMARY_CSV)
    trial_rows = read_csv(FULL_4500_TRIAL_CSV) if FULL_4500_TRIAL_CSV.exists() else []
    comparison_rows = read_csv(COMPARISON_CSV) if COMPARISON_CSV.exists() else []
    tables = compute_tables(summary_rows, comparison_rows, trial_rows)
    generate_4500_figures(summary_rows, tables)
    DOCS.mkdir(exist_ok=True)
    HTML_OUT.write_text(make_html(summary_rows, comparison_rows, tables), encoding="utf-8")
    make_docx(summary_rows, comparison_rows, tables)
    print(f"Wrote {HTML_OUT.relative_to(ROOT)}")
    print(f"Wrote {DOCX_OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
